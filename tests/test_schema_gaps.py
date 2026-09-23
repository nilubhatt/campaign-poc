"""§12.4 — the schema and surface gaps the review named, settled.

The item: "no field today for `asset_link`, `approval_notes`, or market-scoped feedback
patterns. *Decide explicitly whether the tracked client comments 2.5 now ingests ARE the
`approval_notes` gap — in most agency flows a returned deck's comments are exactly that — or
whether a separate structured field is still owed.*"

Thirteen tracker rows accumulated onto it, because every "this needs a field that does not
exist" answer for seven phases was sent here. They are not all the same kind of thing:

  FIELDS      `asset_link`, campaign type (D102), partner (D103), the statuses a campaign can
              actually be in (D38), and retire/revive history (D106).
  SURFACES    attaching a deck to a record that already exists (D39), repairing an asset's
              phase (D123), and the gap that was blocked on the first (D51).
  MEANING     "correction" means three unrelated things (D113).
  DECISIONS   whether tracked comments ARE the approval notes (D15, and D62 behind it), and
              two capability questions where the honest answer is a limit rather than a build
              (D109, D125).

A field added because a review named it, and read by nothing, is the defect this project has
hit most. Every field here is asserted against the thing that READS it.
"""
import pytest

import config
import core
import metrics
import store


def _campaign(conn, title="Colombia", **kw):
    kw.setdefault("market", "LATAM")
    kw.setdefault("status", "concluded")
    kw.setdefault("detail", f"A campaign called {title}, which ran in a market.")
    return core.ingest_campaign(conn, title=title, confirm=True, **kw)["campaign_id"]


# ---------------------------------------------------------------------------------------
# asset_link — the review's first named gap


def test_a_campaign_can_carry_a_link_to_where_the_work_lives(conn):
    """The review's own words. `asset_path` is a file this product STORED; `asset_link` is
    where the work actually lives — the Figma board, the Drive folder, the DAM record — which
    is what somebody reading a judgment six months later needs to go and look at it.

    They are different facts and the product had only the first."""
    cid = _campaign(conn, "Bogota launch",
                    asset_link="https://drive.example.com/folders/bogota-launch")

    assert store.get_campaign(conn, cid)["asset_link"] == (
        "https://drive.example.com/folders/bogota-launch")


def test_the_link_can_be_added_to_a_record_that_already_exists(conn):
    """Most records predate the field, which is the ordinary case for anything added later."""
    cid = _campaign(conn, "Bogota launch")

    core.update_campaign(conn, campaign_id=cid, asset_link="https://figma.example.com/abc")

    assert store.get_campaign(conn, cid)["asset_link"] == "https://figma.example.com/abc"


def test_a_link_that_is_not_a_link_is_refused(conn):
    """It is a pointer somebody will follow. "ask Dana" in that field is a field that looks
    like a link and is not one, and the reader finds out by clicking."""
    with pytest.raises(ValueError, match="link"):
        _campaign(conn, "Bogota launch", asset_link="ask Dana")


def test_a_record_with_no_link_says_nothing_rather_than_empty(conn):
    cid = _campaign(conn, "Bogota launch")

    assert store.get_campaign(conn, cid).get("asset_link") is None


# ---------------------------------------------------------------------------------------
# D38 — the statuses a campaign can actually be in


def test_a_campaign_can_be_cancelled(conn):
    """D38. `cancelled`, `paused` and `on_hold` were REFUSED, with an error explaining there
    was no home for them — and a brief that was cancelled is exactly the record a library
    about what works most needs to keep. "We stopped this one" is an outcome."""
    cid = _campaign(conn, "Bogota launch", status="cancelled")

    assert store.get_campaign(conn, cid)["status"] == "cancelled"


def test_paused_and_on_hold_are_one_state_not_two(conn):
    """They are the same fact in two people's words, and two states nothing distinguishes is
    two checklists, two cells and two gap reports for one situation."""
    first = _campaign(conn, "Bogota launch", status="paused")
    second = _campaign(conn, "Lima launch", status="on hold")

    assert store.get_campaign(conn, first)["status"] == "paused"
    assert store.get_campaign(conn, second)["status"] == "paused"


def test_a_cancelled_campaign_is_not_treated_as_concluded(conn):
    """The distinction that makes the field worth having. A concluded campaign should have
    results and `gaps()` says so when it does not; a cancelled one never will, and reporting
    it as missing its results forever is a gap nobody can close — which is how a library
    teaches people to ignore it."""
    cancelled = _campaign(conn, "Bogota launch", status="cancelled")
    _campaign(conn, "Lima launch", status="concluded")

    reported = core.gaps(conn)

    # The gap's own counted set, not a top-level `campaign_id` — no gap dict has one, so the
    # first version of this assertion compared against `{None}` and could not fail. It is the
    # only test of the behaviour the field exists for, and it was testing nothing.
    outcomes = next(g for g in reported["gaps"] if g["code"] == "few_verified_outcomes")
    asked_for = {a["prefilled_args"].get("campaign_id") for a in outcomes["next_actions"]}
    assert cancelled not in asked_for, \
        "a cancelled campaign is asked for results it will never have"
    assert outcomes["counts"]["campaigns"] == 1, \
        "the cancelled campaign is counted as one that ran"
    # And the coverage table says which it is, rather than calling it not-yet-run: a campaign
    # that was called off has not "not yet" run.
    cells = core.coverage(conn)["cells"]
    assert {c["evidence"] for c in cells if c["stage"] == "cancelled"} == {"never_ran"}


def test_a_cancelled_campaign_is_still_evidence(conn):
    """Kept, not hidden. "We cancelled this" is one of the more useful things a library can
    tell you about a kind of brief, and a record that vanishes from search is one nobody can
    learn from."""
    cid = _campaign(conn, "Bogota launch", status="cancelled",
                    detail="A store opening in Bogota, cancelled when the site fell through.")
    _campaign(conn, "Lima launch", status="concluded", detail="A store opening in Lima.")

    # A FILTERED read. `filter_campaign_ids(conn)` with no arguments returns every row, so the
    # first version of this assertion would have passed with `cancelled` excluded from every
    # filter in the product — it proved only that the row exists.
    assert cid in store.filter_campaign_ids(conn, record_type="campaign")
    assert cid in {c["id"] for c in store.list_campaigns(conn, status="cancelled")}
    # And it is reachable by what it SAYS, which is the half that makes it evidence rather
    # than a row: "we cancelled this" is only useful to somebody searching for store openings.
    found = core.find_similar(conn, text="a store opening in Bogota", top_k=5)
    assert cid in {hit["campaign_id"] for hit in found}


def test_every_status_the_core_accepts_reaches_the_wire(conn):
    """The shape §11.5 hit with a stripped `said_by`, running the other way.

    `store.VALID_STATUSES` grew to five and `mcp_server.Status` still published three, so the
    schema a model reads FIRST said `cancelled` and `paused` were invalid words. A model that
    believes the schema files a cancelled campaign as `concluded` — and `concluded` is what
    every outcome gap, every reconciliation and every calibration figure counts from.

    The same commit that added the statuses carries a comment warning about exactly this. A
    warning is not a guard; this is."""
    import typing

    import mcp_server

    published = next(arg for arg in typing.get_args(mcp_server.Status)
                     if not isinstance(arg, type))

    assert tuple(published.json_schema_extra["enum"]) == store.VALID_STATUSES
    # And the gloss beneath it, which the file itself calls "the copy a model reads first".
    # Read from the SOURCE: a bare string after an `Annotated` assignment is a comment as far
    # as Python is concerned, so `__doc__` is the typing module's and always passes.
    import inspect

    source = inspect.getsource(mcp_server)
    gloss = source.split("Status = _enum(", 1)[1].split('"""')[1]
    for status in store.VALID_STATUSES:
        assert status in gloss, (
            f"{status} is accepted and the model-facing gloss does not mention it"
        )


def test_nothing_still_tells_a_caller_to_leave_a_cancelled_campaign_unset(conn):
    """`enums._EXPLAIN` said "this library has no status for a campaign that was called off —
    leave the status unset", and an unset status is read as `concluded`. The advice steered a
    cancelled campaign into the one bucket it was written to keep it out of. A near-miss table
    that outlives the values it describes is worse than none, because it is confident."""
    import enums

    with pytest.raises(enums.BadValue) as raised:
        enums.normalise("killed", field="status", valid=store.VALID_STATUSES,
                        allow_none=False)

    said = str(raised.value)
    assert "cancelled" in said, "the near-miss has to point at the word that now exists"
    assert "leave the status unset" not in said.lower()
    # And `cancelled` itself is no longer a near-miss at all — it is a value.
    assert enums.normalise("cancelled", field="status", valid=store.VALID_STATUSES,
                           allow_none=False) == "cancelled"
    # Through the store's own normaliser, which is where the synonym table lives — the near-
    # miss explanation only fires for words nothing maps.
    assert store._normalise_status("on hold") == "paused"


# ---------------------------------------------------------------------------------------
# D102 — the checklist is keyed on market, not on what kind of campaign this is


def test_a_campaign_can_say_what_kind_of_campaign_it_is(conn):
    """D102. The review says "the checklist for a campaign TYPE" and renders it "Expected for
    a store launch", and no field carried that: `record_type` is a storage class,
    `collection` is an instance family, and `tags` had no controlled vocabulary until §12.2
    declared one.

    Market was the stand-in, and it is the wrong axis: a store-launch measure was expected of
    every campaign in that market, and of no store launch anywhere else."""
    cid = _campaign(conn, "Bogota launch", campaign_type="store launch")

    assert store.get_campaign(conn, cid)["campaign_type"] == "store launch"


def test_the_type_is_a_declared_vocabulary_not_free_text(conn, tmp_path, monkeypatch):
    """The reason D102 could not be built before §12.2. Free text makes "Store Launch" and
    `store_launch` two checklists — C16's failure — and lets anyone invent a standing
    requirement by typing one."""
    import rulebook

    path = rulebook.overlay_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("version: 'acme-3'\nrules: []\nvocabulary:\n  campaign_types:\n"
                    "    store_launch: ['store launch', 'retail opening']\n", encoding="utf-8")
    rulebook.load.cache_clear()
    try:
        first = _campaign(conn, "Bogota launch", campaign_type="Store Launch")
        second = _campaign(conn, "Lima launch", campaign_type="retail opening")

        assert store.fold_campaign_type(store.get_campaign(conn, first)["campaign_type"]) == \
            store.fold_campaign_type(store.get_campaign(conn, second)["campaign_type"])
    finally:
        rulebook.load.cache_clear()


def test_the_checklist_is_keyed_on_the_campaign_type_it_was_learned_from(conn):
    """D102's actual sentence, and the thing a column on its own does not deliver. The field
    existed, was written by two paths and read by nothing — the defect this project has hit
    most often, and the one the header of this file opens by naming.

    A measure learned entirely from store launches belongs to store launches, in every
    market. Keyed on market it was asked of a seeding brief in Peru that never had footfall
    to uplift, and asked of no store launch in Mexico, which had exactly the same reason to
    report it. Both halves are the same mistake."""
    launches = [_campaign(conn, f"Store launch {n}", market="Peru", campaign_type="store launch",
                          status="concluded") for n in range(3)]
    for n, cid in enumerate(launches):
        metrics.record(conn, campaign_id=cid, key="footfall_uplift_pct", value=4.0 + n,
                       metric_type="actual")
    # Two markets, so the counting gate is satisfied on its own terms and what is left under
    # test is the KEY rather than the threshold.
    fourth = _campaign(conn, "Store launch MX", market="Mexico", campaign_type="store launch",
                       status="concluded")
    metrics.record(conn, campaign_id=fourth, key="footfall_uplift_pct", value=6.0,
                   metric_type="actual")
    metrics.graduate(conn, "footfall_uplift", confirmed_by="R. Vega")

    seeding_in_peru = _campaign(conn, "Peru seeding", market="Peru",
                                campaign_type="influencer seeding", status="concluded")
    launch_in_vietnam = _campaign(conn, "Hanoi launch", market="Vietnam",
                                  campaign_type="store launch", status="concluded")

    assert "footfall_uplift" not in metrics.expected_check(conn, seeding_in_peru)["expected"], (
        "a store-launch measure was asked of every campaign in the market"
    )
    assert "footfall_uplift" in metrics.expected_check(conn, launch_in_vietnam)["expected"], (
        "and of no store launch anywhere else"
    )
    # The sentence a reader sees names the TYPE, which is what tells them why the measure is
    # being asked for — "Expected for a store launch" is the review's own rendering.
    assert "store launch" in metrics.expected_check(conn, launch_in_vietnam)["what_it_means"]


def test_a_measure_learned_across_kinds_of_campaign_stays_keyed_on_market(conn):
    """The other side, and what stops D102's fix narrowing checklists it has no evidence to
    narrow. Mixed evidence says nothing about type; guessing one out of it would remove
    expectations the library never earned the right to remove. `[]` means "keyed on market",
    which is also what every row written before this column existed means."""
    seen_on = [_campaign(conn, "Peru launch", market="Peru", campaign_type="store launch",
                         status="concluded"),
               _campaign(conn, "Peru seeding", market="Peru",
                         campaign_type="influencer seeding", status="concluded"),
               _campaign(conn, "Mexico seeding", market="Mexico",
                         campaign_type="influencer seeding", status="concluded")]
    for cid in seen_on:
        metrics.record(conn, campaign_id=cid, key="reach", value=1000, metric_type="actual")
    metrics.graduate(conn, "reach", confirmed_by="R. Vega")

    anything_in_peru = _campaign(conn, "Peru anything", market="Peru", status="concluded")

    assert "reach" in metrics.expected_check(conn, anything_in_peru)["expected"]


def test_a_type_nobody_declared_is_still_accepted(conn):
    """A declared vocabulary is not a closed one — the same call §12.2 made for markets. A
    customer cannot record the campaign they are running today because they have not got
    round to declaring its type is a gate on doing the work."""
    cid = _campaign(conn, "Bogota launch", campaign_type="pop-up")

    assert store.get_campaign(conn, cid)["campaign_type"] == "pop-up"


# ---------------------------------------------------------------------------------------
# D103 — partners


def test_a_campaign_can_name_the_partner(conn):
    """D103. The review names partners FIRST — "across at least two partners or markets" —
    and nothing in the schema recorded who the partner was, so market stood in for it. Partner
    is arguably the more natural unit of "one partner's house metric"."""
    cid = _campaign(conn, "Bogota launch", partner="Studio Norte")

    assert store.get_campaign(conn, cid)["partner"] == "Studio Norte"


def test_breadth_counts_partners_as_well_as_markets(conn):
    """§8.3's gate is "seen in at least two partners OR markets", and it could only count
    markets. Two campaigns with one partner in two markets is weaker evidence of a general
    rule than two partners in one market — the gate exists to stop one partner's house style
    becoming everybody's standing requirement, and it could not see partners at all."""
    one = _campaign(conn, "Bogota launch", market="Peru", partner="Studio Norte")
    two = _campaign(conn, "Lima launch", market="Peru", partner="Casa Sur")

    assert store.breadth_of(conn, [one, two])["partners"] == 2
    assert store.breadth_of(conn, [one, two])["markets"] == 1


def test_the_gate_lets_two_partners_in_one_market_through(conn):
    """The half §8.3 is NAMED for. `breadth_of` counted partners and `learning.gate` did not
    call it — a function with no caller, which leaves the row closed and the defect exactly
    where it was.

    Two agencies independently writing the same note in one market is real breadth: it is
    precisely NOT one shop's house style, which is what the gate exists to catch. It was
    refused, because only markets could be counted."""
    import corrections as corrections_module

    first = _campaign(conn, "Bogota launch", market="Peru", partner="Studio Norte")
    second = _campaign(conn, "Lima launch", market="Peru", partner="Casa Sur")
    third = _campaign(conn, "Cusco launch", market="Peru", partner="Tercer Piso")
    text = "Creator captions name the product in the first line."
    for cid in (first, second, third):
        noted = corrections_module.note(conn, text=text, campaign_id=cid,
                                        provenance="Client feedback")

    gate = corrections_module.graduation(conn, noted["correction_id"])

    assert gate["partners"] == 3 and gate["markets"] == 1
    assert gate["eligible"], gate["what_it_means"]


def test_a_measure_seen_at_two_partners_in_one_market_can_graduate(conn):
    """The same gate, reached through the OTHER caller. `metrics.graduation` and
    `corrections.graduation` both pass the partner count now and a test of one proves nothing
    about the other — the two-implementations-of-one-rule shape this codebase keeps hitting,
    where the rule is shared and only one route to it was ever exercised."""
    seen_on = [_campaign(conn, "Bogota launch", market="Peru", partner="Studio Norte",
                         status="concluded"),
               _campaign(conn, "Lima launch", market="Peru", partner="Casa Sur",
                         status="concluded"),
               _campaign(conn, "Cusco launch", market="Peru", partner="Tercer Piso",
                         status="concluded")]
    for n, cid in enumerate(seen_on):
        metrics.record(conn, campaign_id=cid, key="footfall_uplift_pct", value=4.0 + n,
                       metric_type="actual")

    gate = metrics.graduation(conn, "footfall_uplift")

    assert gate["partners"] == 3 and gate["markets"] == 1
    assert gate["eligible"], gate["what_it_means"]


def test_one_partner_in_one_market_is_still_refused(conn):
    """And the gate still does the job it was built for. Counting partners is not a loosening
    — three sightings on one agency's decks in one market is the house style the whole
    anti-capture argument is about, and it must not become everybody's requirement."""
    import corrections as corrections_module

    text = "Creator captions name the product in the first line."
    for n in range(3):
        noted = corrections_module.note(
            conn, text=text, provenance="Client feedback",
            campaign_id=_campaign(conn, f"Bogota {n}", market="Peru", partner="Studio Norte"))

    gate = corrections_module.graduation(conn, noted["correction_id"])

    assert not gate["eligible"]
    assert "1 partner" in gate["what_it_means"]


# ---------------------------------------------------------------------------------------
# D106 — a measure that cycles in and out


def test_a_measure_that_cycles_in_and_out_keeps_every_occasion(conn):
    """D106. Only the latest `retired_at` was kept, so a measure that cycled had no record of
    having done so — and "we used to track this" is the whole point of demoting rather than
    deleting. One timestamp cannot say it twice.

    Retirement and revival are not a person's decisions here: the library retires a measure
    nobody has reported for N campaigns and revives one somebody measures again. So the
    history records WHAT the library observed, with `basis: computed` — there is no `said_by`,
    because inventing one would be the library signing its own name to a decision (§11.2)."""
    store.retire_metric(conn, "roas", why="not reported by the last 5 campaigns in Peru")
    store.revive_metric(conn, "roas", why="measured again on the Lima launch")
    store.retire_metric(conn, "roas", why="not reported by the last 5 campaigns in Peru")

    history = store.measure_history(conn, "roas")

    assert [entry["what"] for entry in history] == ["retired", "revived", "retired"]
    assert all(entry["why"] and entry["at"] for entry in history)
    assert all(entry["basis"] == "computed" for entry in history)
    assert not any("said_by" in entry for entry in history), (
        "the library retired it; naming a person would be it signing its own decision"
    )


def test_the_history_says_what_was_observed_each_time(conn):
    """The reason it is worth keeping: a measure dropped because nobody reported it and one
    dropped after a different market stopped is the same word twice with different evidence
    behind it, and the second revival means something the first does not."""
    store.retire_metric(conn, "roas", why="not reported by the last 5 campaigns in Peru")

    entry = store.measure_history(conn, "roas")[0]

    assert entry["why"] == "not reported by the last 5 campaigns in Peru"
    assert entry["what"] == "retired"


def test_the_production_path_fills_the_reason_rather_than_the_default(conn):
    """The test above passes `why` itself, so it proves the column stores what it is given and
    nothing about what real use puts in it. `retire_stale` computed the specific observation
    and passed nothing — every history row in the product read "no longer reported by recent
    campaigns", a sentence true of every retirement and therefore saying nothing about any of
    them. D106 exists so a reader years later can tell one occasion from another."""
    seen_on = [_campaign(conn, "Peru 0", market="Peru", status="concluded"),
               _campaign(conn, "Peru 1", market="Peru", status="concluded"),
               _campaign(conn, "Mexico 0", market="Mexico", status="concluded")]
    for n, cid in enumerate(seen_on):
        metrics.record(conn, campaign_id=cid, key="footfall_uplift_pct", value=4.0 + n,
                       metric_type="actual")
    metrics.graduate(conn, "footfall_uplift", confirmed_by="R. Vega")

    # Campaigns that report SOMETHING and not this one: that is what staleness is, and it
    # takes `learning.RETIREMENT_AFTER` of them. Written against the constant rather than a
    # number, so the test does not quietly stop exercising retirement if the threshold moves.
    import learning

    for n in range(learning.RETIREMENT_AFTER):
        later = _campaign(conn, f"Later {n}", market="Peru", status="concluded")
        metrics.record(conn, campaign_id=later, key="reach", value=1000,
                       metric_type="actual")

    history = store.measure_history(conn, "footfall_uplift")

    retired = [e for e in history if e["what"] == "retired"]
    assert retired, "the measure was never retired, so nothing was observed"
    assert "skipped by" in retired[-1]["why"], retired[-1]
    assert retired[-1]["why"] != "no longer reported by recent campaigns", (
        "the library computed the specific observation and stored the generic one"
    )


def test_retiring_a_measure_that_is_not_expected_records_nothing(conn):
    """`revive_metric` guards on the rowcount and `retire_metric` did not, so retiring a name
    that is not in the registry — or one already retired — appended an occasion anyway. A
    history that records things the library did not do is worse than no history, because the
    whole of D106 is that somebody will read it years later and believe it."""
    store.retire_metric(conn, "never_registered", why="nothing to retire")

    assert store.measure_history(conn, "never_registered") == []

    store.retire_metric(conn, "roas", why="not reported in Peru")
    store.retire_metric(conn, "roas", why="not reported in Peru")

    assert [e["what"] for e in store.measure_history(conn, "roas")] == ["retired"]


def test_a_measures_comings_and_goings_reach_a_surface_somebody_reads(conn):
    """The table was written by two paths and read by nothing: no tool exposed it, and
    `measure_status` — the one surface that answers "where does this measure stand" — did not.
    A history nobody can reach is the stored-and-never-applied defect one table along."""
    store.retire_metric(conn, "roas", why="not reported by the last 5 campaigns in Peru")
    store.revive_metric(conn, "roas", why="measured again on the Lima launch")

    standing = metrics.graduation(conn, "roas")

    assert [e["what"] for e in standing["history"]] == ["retired", "revived"]
    assert "cycles" in standing["history_means"]


def test_the_live_state_still_reads_from_one_place(conn):
    """A history beside a status is two answers to "is this expected". The history is the
    record of occasions; `status` is still what every reader asks."""
    store.retire_metric(conn, "roas", why="not reported by the last 5 campaigns in Peru")
    assert metrics.describe(conn, "roas")["status"] == "retired"

    store.revive_metric(conn, "roas", why="measured again on the Lima launch")

    # `status` is the live answer and the history does not shadow it — the point of the test.
    # The first version asserted only that two rows existed, which is true however `status`
    # behaves: a `revive_metric` that stopped touching `status` altogether stayed green.
    assert metrics.describe(conn, "roas")["status"] == "expected"
    assert [e["what"] for e in store.measure_history(conn, "roas")] == ["retired", "revived"]


# ---------------------------------------------------------------------------------------
# D39 / D51 — attaching a deck to a record that already exists


def test_a_deck_can_be_attached_to_a_record_that_already_exists(conn, tmp_path):
    """D39. `update_campaign` took no `asset_ref` and `upload_image_asset` takes an image, so
    a record whose `commentary_checked` is false could only gain its comments BY BEING
    UPLOADED AGAIN AS A DUPLICATE — which is the offer §5.2 refused in writing.

    Found by 5.2's own test that every offered action must name arguments its tool accepts:
    the obvious offer could not be made, because there was no tool to make it with."""
    from pptx import Presentation

    cid = _campaign(conn, "Bogota launch")
    assert not store.get_campaign(conn, cid)["commentary_checked"]

    deck = tmp_path / "bogota.pptx"
    show = Presentation()
    slide = show.slides.add_slide(show.slide_layouts[5])
    slide.shapes.title.text = "Bogota launch — the store opening deck"
    show.save(deck)

    attached = core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(deck)})

    assert attached["campaign_id"] == cid, "the same record, not a new one"
    assert len(store.filter_campaign_ids(conn)) == 1, "a duplicate is the thing this replaces"
    assert store.get_campaign(conn, cid)["commentary_checked"]


def test_attaching_a_deck_makes_the_record_searchable_by_its_words(conn, tmp_path):
    """A deck attached and not indexed is a file on disk. The point of attaching it is that
    the record can then be found by what the deck says."""
    from pptx import Presentation

    cid = _campaign(conn, "Bogota launch")
    deck = tmp_path / "bogota.pptx"
    show = Presentation()
    slide = show.slides.add_slide(show.slide_layouts[5])
    slide.shapes.title.text = "A flagship store opening with creator seeding in Bogota"
    show.save(deck)

    core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(deck)})

    assert "creator seeding" in " ".join(store.text_on_file(conn, cid)["body"]).lower()


def test_attaching_a_deck_to_a_record_that_has_one_is_refused(conn, tmp_path):
    """Replacing a deck silently would throw away the text every saved judgment was made
    against — §6.3's supersession is how a record gets a new version, and it keeps both."""
    from pptx import Presentation

    deck = tmp_path / "bogota.pptx"
    show = Presentation()
    show.slides.add_slide(show.slide_layouts[5]).shapes.title.text = "The first deck"
    show.save(deck)
    cid = _campaign(conn, "Bogota launch")
    core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(deck)})

    with pytest.raises(ValueError) as raised:
        core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(deck)})

    assert "supersed" in str(raised.value).lower(), "and it names the way to record a new one"


def test_a_file_nothing_can_be_read_from_leaves_the_record_repairable(conn, tmp_path):
    """The refusal above, turned into a trap by the path that writes.

    Attaching a `.png` or a `notes.txt` extracted to nothing, and the record was written
    anyway: empty `deck_text`, an `asset_path` pointing at an unreadable file, the response
    saying "its text is searchable". The record then HAD a deck for every purpose — the real
    `.pptx` was refused when it turned up, and `commentary_never_read` had already stopped
    counting it. A record nobody can repair is worse than one that never got its deck."""
    from PIL import Image

    cid = _campaign(conn, "Bogota launch")
    not_a_deck = tmp_path / "hero.png"
    Image.new("RGB", (32, 32), "navy").save(not_a_deck)

    with pytest.raises(ValueError, match="upload_image_asset"):
        core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(not_a_deck)})

    record = store.get_campaign(conn, cid)
    assert not (record["deck_text"] or "").strip() and not record["asset_path"]

    # And the real deck still goes on afterwards, which is the whole point of refusing.
    from pptx import Presentation

    deck = tmp_path / "bogota.pptx"
    show = Presentation()
    show.slides.add_slide(show.slide_layouts[5]).shapes.title.text = "The real deck"
    show.save(deck)

    core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(deck)})

    assert "real deck" in store.get_campaign(conn, cid)["deck_text"]


def test_a_truncated_file_leaves_neither_a_row_nor_an_orphan(conn, tmp_path):
    """The same trap through the other door. A zero-byte `.pptx` raised out of extraction —
    after the file had already been copied into the asset store, so every attempt left a
    stray file behind. Extraction comes first now: nothing is kept until something was read."""
    cid = _campaign(conn, "Bogota launch")
    truncated = tmp_path / "half.pptx"
    truncated.write_bytes(b"")
    before = len(list(config.ASSET_DIR.glob("*"))) if config.ASSET_DIR.exists() else 0

    with pytest.raises(ValueError, match="could not be read as a deck"):
        core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(truncated)})

    after = len(list(config.ASSET_DIR.glob("*"))) if config.ASSET_DIR.exists() else 0
    assert after == before, "a refused attach left a file in the asset store"
    assert not store.get_campaign(conn, cid)["asset_path"]


def test_two_decks_arriving_at_once_cannot_both_be_attached(conn):
    """A check-then-act, and the reason the condition lives in the WHERE clause.

    `core.attach_deck` checks "does this record have a deck" and must, because it owes a
    sentence about supersession rather than a bare False. But two attaches racing both passed
    that check, and the row ended up holding one deck's TEXT with the other deck's COMMENTS —
    a record that says one thing and is annotated with remarks about another, which is the
    misattribution this whole product is built against.

    Reached through the store directly, because that is where a race arrives: the caller-side
    check is what a second caller has already passed by the time it matters."""
    cid = _campaign(conn, "Bogota launch")

    assert store.attach_deck_to_campaign(conn, cid, deck_text="Deck A", asset_path="a.pptx",
                                         commentary_checked=True)
    assert not store.attach_deck_to_campaign(conn, cid, deck_text="Deck B",
                                             asset_path="b.pptx", commentary_checked=True), (
        "the second attach overwrote the first, so the text and the comments disagree"
    )

    record = store.get_campaign(conn, cid)
    assert record["deck_text"] == "Deck A" and record["asset_path"] == "a.pptx"


def test_an_attached_deck_is_read_for_promises_like_any_other(conn, tmp_path):
    """The record repaired exactly as `commentary_never_read` recommends had no §9.3
    commitments at all, so `check_commitments` and the briefed half of `compare_execution`
    read empty on precisely the records the product had just told somebody to fix.

    Both other write paths extract them. The edit path's own comment calls omitting this "the
    'learning from the wrong document' failure this phase exists to stop"."""
    from pptx import Presentation

    import commitments

    cid = _campaign(conn, "Bogota launch")
    deck = tmp_path / "bogota.pptx"
    show = Presentation()
    slide = show.slides.add_slide(show.slide_layouts[1])
    slide.shapes.title.text = "Bogota launch"
    slide.placeholders[1].text = ("We will seed one colourway per creator.\n"
                                  "Every asset will carry the currency.")
    show.save(deck)

    core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(deck)})

    promised = commitments.summary_for(conn, cid)
    assert promised.get("count", 0) >= 1, promised
    assert any("colourway" in p["text"] for p in promised["promises"]), promised


def test_a_record_whose_deck_was_never_read_is_reported_as_a_gap(conn):
    """D51, which was blocked on D39: "the only offer available today creates a duplicate
    campaign, which is the offer 5.2 refused in writing". A gap whose only remedy is a thing
    the product refuses to do is not a gap worth reporting.

    Now there is a remedy, so the gap can be."""
    _campaign(conn, "Bogota launch")
    _campaign(conn, "Lima launch")

    codes = {gap["code"] for gap in core.gaps(conn)["gaps"]}

    assert "commentary_never_read" in codes


def test_the_gap_offers_the_tool_that_fixes_it_rather_than_a_duplicate(conn):
    """The reason D51 waited. An offer that creates a duplicate record is worse than no offer:
    §5.2 refused it in writing, and a gap that recommends it teaches somebody to corrupt their
    own library."""
    _campaign(conn, "Bogota launch")
    _campaign(conn, "Lima launch")

    gap = next(g for g in core.gaps(conn)["gaps"] if g["code"] == "commentary_never_read")
    offered = {a["tool"] for a in gap.get("next_actions", [])}

    assert "attach_deck" in offered
    assert "upload_campaign" not in offered, "that is the duplicate 5.2 refused"


# ---------------------------------------------------------------------------------------
# D123 — repairing an asset's phase


def test_an_assets_phase_can_be_corrected(conn, tmp_path):
    """D123. "A model that guessed `proposed` on fourteen photographs has produced an
    unrepairable record." `phase` decides whether an image is BRIEFED creative or what
    actually ran — §9.1 says everything falls out of that — and there was no setter."""
    from PIL import Image

    cid = _campaign(conn, "Bogota launch")
    path = tmp_path / "shot.png"
    Image.new("RGB", (64, 64), "navy").save(path)
    asset = core.ingest_image_asset(conn, campaign_id=cid,
                                    asset_ref={"path": str(path)}, phase="proposed")

    core.update_asset(conn, asset_id=asset["asset_id"], phase="delivered",
                      why="These are the event photographs, not the boards",
                      said_by="R. Vega")

    assert store.get_asset(conn, asset["asset_id"])["phase"] == "delivered"
    # KEPT, not merely returned. `why` was required, echoed back and stored nowhere, so the
    # sentence this tool calls "the part a reader needs six months later" was gone the moment
    # the call returned — and a mutation deleting the `why` guard left every test green.
    wrote = store.authorship_for(conn, "asset", asset["asset_id"])
    assert wrote["on_behalf_of"]["name"] == "R. Vega"
    assert wrote["on_behalf_of"]["why"] == "These are the event photographs, not the boards"


def test_correcting_a_phase_records_who_said_so_and_why(conn, tmp_path):
    """It changes what the image IS evidence of, so it is a consequential write — and every
    consequential write in this product takes a reason and a person (§11.2)."""
    from PIL import Image

    cid = _campaign(conn, "Bogota launch")
    path = tmp_path / "shot.png"
    Image.new("RGB", (64, 64), "navy").save(path)
    asset = core.ingest_image_asset(conn, campaign_id=cid,
                                    asset_ref={"path": str(path)}, phase="proposed")

    with pytest.raises(ValueError, match="PERSON"):
        core.update_asset(conn, asset_id=asset["asset_id"], phase="delivered",
                          why="They are photographs", said_by="the system")

    assert store.get_asset(conn, asset["asset_id"])["phase"] == "proposed", (
        "and nothing was written before the refusal"
    )


def test_a_record_that_never_had_a_deck_can_say_so(conn):
    """The other way D51's gap closes, and without it the gap is the thing it was avoiding.

    Some decks are genuinely lost — the agency that made them no longer exists — so "attach
    the deck" is an instruction to produce something nobody has. Reported forever, that is the
    gap everybody learns to ignore, which is why it went unreported for seven phases.

    The first version of this test set the gap aside on a PASTED-TEXT brief, which the gap no
    longer reports at all and which `attach_deck` refuses outright: it demonstrated the
    product recommending an action its own tool rejects and then silencing the complaint. The
    record here has no deck, which is the population the offer actually fits."""
    core.ingest_campaign(conn, title="A deck nobody can find any more",
                         detail="What we remember of it.", confirm=True)
    assert "commentary_never_read" in {g["code"] for g in core.gaps(conn)["gaps"]}

    core.answer_gap(conn, code="commentary_never_read", answer="not_applicable",
                    note="That agency is gone and the files went with it.",
                    said_by="R. Vega")

    assert "commentary_never_read" not in {g["code"] for g in core.gaps(conn)["gaps"]}


def test_setting_a_gap_aside_does_not_silence_it_for_records_added_later(conn):
    """§10.2's set-aside is keyed on the gap's CODE, and the evidence under a code moves.

    "Those decks are gone" is a statement about what was on file when somebody said it. Taken
    as a permanent instruction, one decision silenced the gap for every record uploaded
    afterwards — the product agreeing to stop mentioning something it had not yet seen, which
    is the same shape as a judgment that cannot be falsified. It comes back, and says why."""
    import time as _time

    core.ingest_campaign(conn, title="Lost to the old agency", detail="From memory.",
                         confirm=True)
    core.answer_gap(conn, code="commentary_never_read", answer="not_applicable",
                    note="That agency is gone.", said_by="R. Vega")
    assert "commentary_never_read" not in {g["code"] for g in core.gaps(conn)["gaps"]}

    _time.sleep(0.01)
    core.ingest_campaign(conn, title="Uploaded this morning", detail="Deck to follow.",
                         confirm=True)

    gap = next(g for g in core.gaps(conn)["gaps"] if g["code"] == "commentary_never_read")
    # Reported again — and the note somebody wrote travels with it, so nobody has to work out
    # whether this was already looked at.
    assert gap["reopened"]["what_it_means"]
    assert gap["answered"]["said_by"] == "R. Vega"


# ---------------------------------------------------------------------------------------
# D15 / D62 — the decision the item asks for, in as many words
#
# "Decide explicitly whether the tracked client comments 2.5 now ingests ARE the
# `approval_notes` gap — in most agency flows a returned deck's comments are exactly that —
# or whether a separate structured field is still owed."
#
# THEY ARE. A returned deck's tracked comments are the approval notes: they are what the
# client wrote when they sent it back. And they arrive with more than a free-text field could
# ever hold — an author, an anchor to the slide, a date, a reply thread — because they came
# out of the file rather than being retyped into a box.
#
# A separate `approval_notes` field would be a second place for the same thing, which is this
# codebase's signature defect, and the worse copy: retyped, unanchored, unattributed. What was
# genuinely missing is not the notes but the VERDICT — whether the deck came back approved —
# and that is one field rather than a parallel store.


def test_a_returned_decks_comments_are_the_approval_notes(conn, tmp_path):
    """The decision, asserted against the thing that reads them. The comments carry an author
    and an anchor, which is what makes them worth more than a text box."""
    from pptx import Presentation
    from pptx.util import Inches

    cid = _campaign(conn, "Bogota launch")
    deck = tmp_path / "returned.pptx"
    show = Presentation()
    slide = show.slides.add_slide(show.slide_layouts[5])
    slide.shapes.title.text = "Bogota launch"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "The creative is approved; the timing is not."
    slide.notes_slide.notes_text_frame.text = "Client: approved subject to the date moving."
    show.save(deck)

    core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(deck)})

    # `get_commentary`, because the claim is that these carry MORE than a free-text box: an
    # author, an anchor, a kind. `text_on_file` returns bare strings, so asserting on it
    # proved only that something was stored — which a retyped `approval_notes` box would also
    # satisfy, and that box is the thing this decision rejects.
    commentary = store.get_commentary(conn, cid)
    assert commentary, "the returned deck's notes are the approval notes and must be readable"
    assert all(c.get("kind") for c in commentary), "each says what kind of remark it is"
    assert any(c.get("anchor") or c.get("slide") or c.get("page") for c in commentary), (
        "an anchor is half of what makes a comment worth keeping, and no retyped box has one"
    )


def test_the_record_can_say_whether_it_came_back_approved(conn):
    """What was actually missing. The NOTES are the comments; the VERDICT is not in them —
    "the timing is not" is a note, and whether the deck was signed off is a different fact
    that a reader six months later needs first."""
    cid = _campaign(conn, "Bogota launch")

    core.update_campaign(conn, campaign_id=cid, approval="approved_with_changes",
                         approval_note="Signed off on the creative, not the dates.",
                         said_by="R. Vega")

    record = store.get_campaign(conn, cid)
    assert record["approval"] == "approved_with_changes"
    assert record["approval_note"] == "Signed off on the creative, not the dates."
    # The NAME is stored, not just validated. Nothing asserted this, so a version that
    # checked `said_by` was a person and then dropped it stayed green — leaving the library
    # asserting the deck was signed off with no way to say by whom, which is the one thing
    # §11.2 exists to prevent.
    assert record["approval_by"] == "R. Vega"


def test_the_approval_says_who_said_so(conn):
    """§11.2 everywhere it applies: "the client approved it" with nobody's name against it is
    an opinion the library holds and cannot attribute."""
    cid = _campaign(conn, "Bogota launch")

    with pytest.raises(ValueError, match="PERSON"):
        core.update_campaign(conn, campaign_id=cid, approval="approved",
                             approval_note="Signed off.", said_by="the client")

    # And omitting the name altogether is refused too. Validating `said_by` when it happens
    # to arrive is not the same rule as requiring it: the gap between the two is a row saying
    # `approval: approved` with `approval_by` empty.
    with pytest.raises(ValueError, match="said_by"):
        core.update_campaign(conn, campaign_id=cid, approval="approved")

    assert store.get_campaign(conn, cid)["approval"] is None


def test_there_is_no_second_place_for_the_notes_themselves(conn):
    """The decision's other half, which is what stops it becoming the defect it avoids. There
    is no `approval_notes` free-text field: the notes ARE the commentary, which is anchored
    and attributed, and a parallel box would be a retyped copy of the same thing — worse in
    every way and drifting from the first day."""
    columns = {row[1] for row in conn.execute("PRAGMA table_info(campaigns)")}

    assert "approval_notes" not in columns
    assert "approval" in columns and "approval_note" in columns


def _deck_a_reviewer_wrote_on(path, remark, author="R. Vega"):
    """A deck carrying a remark SOMEBODY ELSE wrote on it — a PDF `/Text` annotation.

    Not a speaker note. The first version of the two tests below used `notes_slide`, which is
    the deck author's own presenter script, and asserted that the product reported it as "what
    the client wrote when they sent it back". They passed, because the function under test
    read a list with the `kind` stripped off and could not tell the two apart — so the test
    that existed to prove attribution was honest was demonstrating the misattribution. Review
    caught it; this builds the thing the test was always describing.
    """
    from pypdf import PdfWriter
    from pypdf.annotations import Text
    from pypdf.generic import NameObject, TextStringObject

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_annotation(page_number=0, annotation=Text(text=remark, rect=(1, 1, 20, 20)))
    target = writer.pages[0]["/Annots"][-1].get_object()
    target[NameObject("/T")] = TextStringObject(author)
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def _a_remark_on_the_deck(conn, campaign_id, text, author="R. Vega"):
    """One commentary chunk on a record that already has a body, stored the way ingestion
    stores one: its own kind, its own author, its own anchor, never packed with another.

    `attach_deck` cannot be used here — it refuses a record that already has a deck, which is
    the behaviour two rows up — and the body is what these tests need on the other layer.
    """
    store.insert_chunks(conn, campaign_id, [text], kind="commentary",
                        sources=[{"kind": "comment", "author": author, "anchor": "slide 2"}])
    return text


def test_the_earlier_decks_comments_count_as_corrections_taken(conn, tmp_path):
    """D62. "In the Colombia case 'corrections' meant the client's tracked comments at least
    as much as the library's findings." A diff that reports only which of ITS OWN findings
    were addressed is reporting on itself; what the client actually asked for is in the deck
    they sent back, and it was not being read as anything."""
    first = _campaign(conn, "Bogota launch v1")
    deck = _deck_a_reviewer_wrote_on(tmp_path / "v1.pdf",
                                     "Client: drop the third colourway.")
    core.attach_deck(conn, campaign_id=first, asset_ref={"path": str(deck)})

    second = core.ingest_campaign(
        conn, title="Bogota launch v2", market="LATAM", status="proposed",
        detail="One colourway per creator, as asked.", supersedes=first,
        confirm=True)["campaign_id"]

    diff = core.diff_campaigns(conn, earlier=first, later=second)

    asked = diff["what_the_client_asked_for"]
    assert asked["count"] == 1, asked
    assert "colourway" in " ".join(item["text"] for item in asked["items"]).lower()
    # The AUTHOR travels with the ask. "Who wanted this" is half of what makes an ask worth
    # answering, and the first version reported the text alone — so an unattributed remark and
    # one a named client wrote read identically.
    assert asked["items"][0]["author"] == "R. Vega"
    assert asked["authors"] == ["R. Vega"] and asked["some_unattributed"] is False


def test_the_verdict_changes_what_the_asks_mean(conn, tmp_path):
    """D15's verdict read beside D62's asks, which is where storing it earns its place.

    "Drop the third colourway" under `rejected` is why the deck did not proceed; the same
    sentence under `approved_with_changes` is what the sign-off was conditional on; under
    `approved` it is a remark made in passing. The field was written by two paths and read by
    nothing — and this is the surface where not having it changes what a quotation means."""
    first = _campaign(conn, "Bogota launch v1")
    core.attach_deck(conn, campaign_id=first, asset_ref={"path": str(
        _deck_a_reviewer_wrote_on(tmp_path / "v1.pdf", "Client: drop the third colourway."))})
    core.update_campaign(conn, campaign_id=first, approval="rejected",
                         approval_note="Not in this form.", said_by="R. Vega")
    second = core.ingest_campaign(conn, title="Bogota launch v2", market="LATAM",
                                  status="proposed", detail="One colourway.",
                                  supersedes=first, confirm=True)["campaign_id"]

    asked = core.diff_campaigns(conn, earlier=first,
                                later=second)["what_the_client_asked_for"]

    assert asked["approval"] == "rejected" and asked["approval_by"] == "R. Vega"
    assert "did not survive" in asked["approval_means"]


def test_a_deck_with_no_verdict_on_it_says_nothing_rather_than_guessing(conn, tmp_path):
    """An unset `approval` means nobody has recorded a decision, and the four values are all
    somebody's act. Reporting a default here would be the library inventing a verdict on a
    deck to say something about the remarks on it."""
    first = _campaign(conn, "Lima launch v1")
    core.attach_deck(conn, campaign_id=first, asset_ref={"path": str(
        _deck_a_reviewer_wrote_on(tmp_path / "v1.pdf", "Client: tighten the copy."))})
    second = core.ingest_campaign(conn, title="Lima launch v2", market="LATAM",
                                  status="proposed", detail="Tightened.", supersedes=first,
                                  confirm=True)["campaign_id"]

    asked = core.diff_campaigns(conn, earlier=first,
                                later=second)["what_the_client_asked_for"]

    assert "approval" not in asked and "approval_means" not in asked
    assert asked["count"] == 1


def test_a_measure_retired_before_the_history_existed_does_not_read_as_revived(conn):
    """An upgrade cannot invent an occasion, and it must not leave one out either.

    A measure retired by a release that predates `measure_history` has no `retired` row, so
    the first thing its history said after the upgrade was "revived" — a record reading as
    though the library brought back something it had never demoted. The retirement is
    derivable: `retired_at` is on the row and is when it happened."""
    store.register_metric(conn, canonical="legacy_measure", display_name="Legacy",
                          unit=None, direction=None, aliases=[])
    conn.execute("UPDATE metric_registry SET status = 'retired', retired_at = ? "
                 "WHERE canonical = ?", (1_700_000_000.0, "legacy_measure"))
    conn.execute("DELETE FROM measure_history WHERE canonical = 'legacy_measure'")
    conn.commit()

    store._migrate_schema(conn)

    history = store.measure_history(conn, "legacy_measure")
    assert [e["what"] for e in history] == ["retired"]
    assert history[0]["at"] == 1_700_000_000.0

    # And running the migration twice does not write it twice.
    store._migrate_schema(conn)
    assert len(store.measure_history(conn, "legacy_measure")) == 1


def test_the_agencys_own_speaker_notes_are_not_reported_as_the_clients_asks(conn, tmp_path):
    """The notes slide is the deck author talking to themselves. Reported as "what the client
    wrote when they sent it back", "Presenter: remember to smile, and skip slide 4" becomes an
    instruction from the client — a false attribution produced by a product whose whole thesis
    is that it can say where a claim came from. The three commentary kinds are not the same
    claim and the function has to read the kind."""
    from pptx import Presentation

    first = _campaign(conn, "Lima launch v1")
    deck = tmp_path / "v1.pptx"
    show = Presentation()
    slide = show.slides.add_slide(show.slide_layouts[5])
    slide.shapes.title.text = "Lima launch"
    show.save(deck)
    slide.notes_slide.notes_text_frame.text = "Presenter: remember to smile, skip slide 4."
    show.save(deck)
    core.attach_deck(conn, campaign_id=first, asset_ref={"path": str(deck)})

    second = core.ingest_campaign(conn, title="Lima launch v2", market="LATAM",
                                  status="proposed", detail="Tightened.", supersedes=first,
                                  confirm=True)["campaign_id"]

    asked = core.diff_campaigns(conn, earlier=first, later=second)["what_the_client_asked_for"]

    assert asked["count"] == 0 and asked["items"] == []
    # And it says WHY it is empty, because "nobody commented" and "the only remarks were the
    # deck's own script" are different answers and a reader acts differently on each.
    assert "speaker note" in asked["what_it_means"]
    # The note is still STORED and searchable — it is dropped from this surface, not from the
    # library. A deck's own notes are often the only place the brief explains itself.
    assert any(c["kind"] == "speaker_note" for c in store.get_commentary(conn, first))


def test_the_clients_asks_are_not_reported_as_the_librarys_findings(conn, tmp_path):
    """They are different claims and a reader has to be able to tell them apart: one is what
    the client wrote on the deck, the other is what this product decided. Blending them would
    let the library take credit for the client's own instruction."""
    from pptx import Presentation

    first = _campaign(conn, "Bogota launch v1")
    deck = tmp_path / "v1.pptx"
    show = Presentation()
    show.slides.add_slide(show.slide_layouts[5]).shapes.title.text = "Bogota launch"
    show.save(deck)
    core.attach_deck(conn, campaign_id=first, asset_ref={"path": str(deck)})
    second = core.ingest_campaign(conn, title="Bogota launch v2", market="LATAM",
                                  status="proposed", detail="A revision.", supersedes=first,
                                  confirm=True)["campaign_id"]

    diff = core.diff_campaigns(conn, earlier=first, later=second)

    assert diff["what_the_client_asked_for"]["basis"] == "computed"
    assert "adopted" in diff and diff["what_the_client_asked_for"] is not diff["adopted"]


# ---------------------------------------------------------------------------------------
# D63 — what the two decks actually say differently


def test_the_diff_says_which_passages_were_added_and_removed(conn, tmp_path):
    """D63. "Added and removed passages between two versions, by layer, from the chunks
    already stored per version."

    `fact_changes` says a budget moved and `record_changes` says a field did; neither says
    the third colourway paragraph is gone, which is what somebody comparing two decks is
    actually looking at. The chunks are already stored per version — nothing had read them
    against each other."""
    first = core.ingest_campaign(
        conn, title="Bogota v1", market="LATAM", status="proposed", confirm=True,
        deck_text="Seed four colourways per creator.\n\nPaid social behind the launch."
    )["campaign_id"]
    second = core.ingest_campaign(
        conn, title="Bogota v2", market="LATAM", status="proposed", supersedes=first,
        confirm=True,
        deck_text="Seed one colourway per creator.\n\nPaid social behind the launch."
    )["campaign_id"]

    passages = core.diff_campaigns(conn, earlier=first, later=second)["passages"]

    assert any("four colourways" in p["text"] for p in passages["removed"])
    assert any("one colourway" in p["text"] for p in passages["added"])
    assert not any("Paid social" in p["text"] for p in passages["added"] + passages["removed"])


def test_the_passage_diff_says_which_layer_each_came_from(conn, tmp_path):
    """"by layer" is the item's own word and it is the load-bearing part: a paragraph removed
    from the BODY is the brief changing, and a comment that is gone is the client's remark
    being dropped. Reporting them as one list would make a deleted objection look like an
    edit."""
    first = core.ingest_campaign(
        conn, title="Bogota v1", market="LATAM", status="proposed", confirm=True,
        deck_text="Seed four colourways per creator.")["campaign_id"]
    second = core.ingest_campaign(
        conn, title="Bogota v2", market="LATAM", status="proposed", supersedes=first,
        confirm=True, deck_text="Seed one colourway per creator.")["campaign_id"]

    _a_remark_on_the_deck(conn, first, "Cut the fourth colourway.")

    passages = core.diff_campaigns(conn, earlier=first, later=second)["passages"]

    # BOTH layers, each on the right passage. Asserting only that the label is one of two
    # words passed while every passage was mislabelled `commentary` — the field existed and
    # said nothing, which is worse than no field, because a deleted client objection reading
    # as an edit is exactly what the layer is for.
    where = {p["text"]: p["layer"] for p in passages["removed"]}
    assert any("colourways per creator" in t and layer == "body"
               for t, layer in where.items()), where
    assert any("fourth colourway" in t and layer == "commentary"
               for t, layer in where.items()), where


def test_a_dropped_client_comment_is_not_crowded_out_by_a_rewritten_body(conn, tmp_path):
    """The layer rule at the point where it costs something.

    There is always more body than commentary, so a body-first cap meant a deck with fifteen
    rewritten paragraphs and two dropped client comments returned ten body passages and ZERO
    commentary, both comments folded into an undifferentiated `more`. The layer this function
    calls "usually the more interesting of the two" was the one truncation removed, every
    time — a bound that inverts the priority it was written to protect."""
    first = core.ingest_campaign(
        conn, title="Bogota v1", market="LATAM", status="proposed", confirm=True,
        deck_text="\n\n".join(f"Paragraph {n} of the brief." for n in range(15))
    )["campaign_id"]
    _a_remark_on_the_deck(conn, first, "Drop the third colourway entirely.")
    second = core.ingest_campaign(
        conn, title="Bogota v2", market="LATAM", status="proposed", supersedes=first,
        confirm=True,
        deck_text="\n\n".join(f"Something else, {n}." for n in range(15))
    )["campaign_id"]

    removed = core.diff_campaigns(conn, earlier=first, later=second)["passages"]["removed"]

    assert any(p["layer"] == "commentary" for p in removed), (
        "the client's dropped remark was pushed out of the reply by body paragraphs"
    )


def test_the_same_words_in_two_unicode_forms_are_not_a_change(conn):
    """"Café" is one codepoint on one machine and "e" plus a combining accent on another. The
    two are the same sentence to every reader and different bytes to Python, so a deck
    round-tripped through a Mac reported every accented paragraph as both added and removed —
    the diff saying the brief changed when nobody touched it."""
    import unicodedata

    line = "Café culture drives the Bogotá launch."
    first = core.ingest_campaign(
        conn, title="Bogota v1", market="LATAM", status="proposed", confirm=True,
        deck_text=unicodedata.normalize("NFC", line))["campaign_id"]
    second = core.ingest_campaign(
        conn, title="Bogota v2", market="LATAM", status="proposed", supersedes=first,
        confirm=True, deck_text=unicodedata.normalize("NFD", line))["campaign_id"]

    passages = core.diff_campaigns(conn, earlier=first, later=second)["passages"]

    assert passages["added"] == [] and passages["removed"] == []


def test_two_identical_decks_produce_no_passage_noise(conn):
    """A diff that reports every paragraph as changed because the chunker packed them
    differently is one nobody reads twice."""
    first = core.ingest_campaign(
        conn, title="Bogota v1", market="LATAM", status="proposed", confirm=True,
        deck_text="Seed one colourway per creator.\n\nPaid social behind it.")["campaign_id"]
    second = core.ingest_campaign(
        conn, title="Bogota v2", market="LATAM", status="proposed", supersedes=first,
        confirm=True,
        deck_text="Seed one colourway per creator.\n\nPaid social behind it.")["campaign_id"]

    passages = core.diff_campaigns(conn, earlier=first, later=second)["passages"]

    assert passages["added"] == [] and passages["removed"] == []


def test_the_passage_diff_is_bounded(conn):
    """Two long decks with nothing in common would otherwise return both decks in full, in a
    reply somebody has to read."""
    first = core.ingest_campaign(
        conn, title="Bogota v1", market="LATAM", status="proposed", confirm=True,
        deck_text="\n\n".join(f"Paragraph {n} about the launch." for n in range(60))
    )["campaign_id"]
    second = core.ingest_campaign(
        conn, title="Bogota v2", market="LATAM", status="proposed", supersedes=first,
        confirm=True,
        deck_text="\n\n".join(f"Something else entirely, {n}." for n in range(60))
    )["campaign_id"]

    passages = core.diff_campaigns(conn, earlier=first, later=second)["passages"]

    assert len(passages["added"]) <= 10 and len(passages["removed"]) <= 10
    assert passages["more"] > 0, "and it says how many it did not show"


# ---------------------------------------------------------------------------------------
# D113 — "correction" means three unrelated things


def test_each_kind_of_correction_says_which_kind_it_is(conn):
    """D113. The word means three different things in this product:

      1. `diff_campaigns`' corrections TAKEN — findings the next version addressed (§5.4/6.3)
      2. a `recomputed` metric VALUE — a number that was wrong and was fixed (§8.1)
      3. a STANDING correction — a rule the library learned from repeated feedback (§8.6)

    A reader who sees "3 corrections" cannot tell which, and the three carry completely
    different weight: one is a diff observation, one is a data fix, one is a rule that judges
    every future brief in a market.

    Renaming the stored vocabulary of two of them would break every saved row, so what is
    asserted here is that each SURFACE says which kind it means, in the words a reader uses."""
    cid = _campaign(conn, "Bogota launch")
    import corrections as corrections_module

    noted = corrections_module.note(conn, text="Never put the logo on a dark background.",
                                    campaign_id=cid, provenance="Client call, 3 March")

    described = corrections_module.describe(conn, noted["correction_id"])
    assert "standing" in described["what_it_means"].lower()


def test_the_diffs_corrections_are_not_called_standing_corrections(conn):
    """The one most likely to mislead: a diff reporting "corrections adopted" beside a
    judgment that cites `standing_corrections` puts two different things under one word, on
    one screen."""
    first = core.ingest_campaign(conn, title="Bogota v1", market="LATAM", status="proposed",
                                 confirm=True, deck_text="Four colourways.")["campaign_id"]
    second = core.ingest_campaign(conn, title="Bogota v2", market="LATAM", status="proposed",
                                  supersedes=first, confirm=True,
                                  deck_text="One colourway.")["campaign_id"]

    diff = core.diff_campaigns(conn, earlier=first, later=second)

    # The surface SAYS which kind it means. Asserting only that a phrase is absent passed on
    # a response that said nothing at all — `diff_campaigns` has no `what_it_means` key, so
    # the check ran against five integers and could not fail.
    said = diff["what_a_correction_means_here"]
    assert "observation about two documents" in said.lower()
    assert "not a STANDING correction" in said, (
        "it has to name what it is NOT, or the reader has to already know the distinction"
    )


def test_a_recomputed_metric_is_not_called_a_correction(conn):
    """The third. §8.1's `recomputed` is a NUMBER that was wrong — nothing to do with a rule
    or a finding — and calling it a correction puts a data fix in the same vocabulary as a
    standing requirement."""
    import metrics

    cid = _campaign(conn, "Bogota launch")
    metrics.record(conn, campaign_id=cid, key="roas", value=3.4, metric_type="actual")
    again = metrics.record(conn, campaign_id=cid, key="roas", value=4.1,
                           metric_type="actual")

    # It says which kind of correction a recomputed value IS, rather than avoiding the word.
    # The first version asserted the word was absent from the response — which it was, because
    # the response is four keys and none of them says anything — so the test passed on a
    # surface that had never been given the sentence it was supposed to carry.
    # `_strip_decoration` reads the source OUT of the key — that is §8.1's whole design, and
    # `planned_roas_july_recomputed` is the review's own example of it.
    corrected = metrics.record(conn, campaign_id=cid, key="roas_recomputed", value=4.1,
                               metric_type="actual")
    said = corrected["what_a_correction_means_here"]
    assert "corrected number" in said.lower()
    assert "not a STANDING correction" in said
    # And an ordinary stated value does not carry it: a sentence on every write is one nobody
    # reads, and this one is only true of a recomputation.
    assert "what_a_correction_means_here" not in again


# ---------------------------------------------------------------------------------------
# D125 — a promise about a COUNT
#
# "Single colourway: not observed, four colourways present" is the review's sentence, and the
# product produces only the first half. The row settles it: refusing is RIGHT, because a
# max-similarity search confirms a universal claim by construction — for "a single colourway",
# four colourways in frame make the match stronger, not weaker.
#
# What §12.4 settles is the other half: counting what IS in a set of photographs is a
# different INSTRUMENT — object detection and attribute extraction, not retrieval — and this
# product ships a retriever and a perceptual hash. That is an accepted limit rather than work
# owed, and the thing that makes it honest is the product saying which instrument would
# answer, so a reader knows this is a boundary and not a bug.


def test_a_promise_about_a_count_says_what_would_answer_it(conn, tmp_path):
    """The refusal already says why a similarity search cannot answer it. What it did not say
    is what COULD — so a reader cannot tell a boundary of this product from something that is
    broken or unfinished, and the difference decides whether they wait or go and look."""
    from PIL import Image

    cid = _campaign(conn, "Bogota launch")
    import commitments

    commitments.add(conn, campaign_id=cid, text="A single colourway per creator.",
                    source_line="Seeding: a single colourway per creator.")
    path = tmp_path / "shot.png"
    Image.new("RGB", (64, 64), "navy").save(path)
    core.ingest_image_asset(conn, campaign_id=cid, asset_ref={"path": str(path)},
                            phase="delivered")

    checked = commitments.check(conn, campaign_id=cid)
    item = next(i for i in checked["items"] if "colourway" in i["text"])

    assert item["verdict"] == "unchecked"
    said = item["what_it_means"].lower()
    assert "count" in said or "how many" in said
    assert "detect" in said or "instrument" in said, (
        "it has to name what would answer it, or a limit reads as a defect"
    )


def test_the_limit_is_not_dressed_up_as_a_result(conn, tmp_path):
    """The failure the refusal exists to prevent, still held: `unchecked` is never a pass, and
    a count promise must not come back `present` because a photograph resembled the phrase."""
    from PIL import Image

    cid = _campaign(conn, "Bogota launch")
    import commitments

    commitments.add(conn, campaign_id=cid, text="Only one hero product in frame.",
                    source_line="Creative: only one hero product in frame.")
    commitments.add(conn, campaign_id=cid, text="The hero shot is navy.",
                    source_line="Creative: the hero shot is navy.")
    path = tmp_path / "shot.png"
    Image.new("RGB", (64, 64), "navy").save(path)
    core.ingest_image_asset(conn, campaign_id=cid, asset_ref={"path": str(path)},
                            phase="delivered")

    item = next(i for i in commitments.check(conn, campaign_id=cid)["items"]
                if "hero product" in i["text"])

    assert item["verdict"] == "unchecked"
    assert item["similarity"] is None, "a similarity beside an unchecked verdict reads as one"
    # `_unchecked` hardcodes `similarity: None` for every reason it refuses, so the assert
    # above is asserting a constant and would hold with the count rule deleted entirely. What
    # distinguishes this refusal is the REASON, and an ordinary promise on the same record
    # gives a different one: this is about what a max-similarity search cannot establish, not
    # about whether the photographs happen to be indexed today. Collapse the two sentences
    # into one and a reader can no longer tell a permanent boundary from a missing model.
    ordinary = next(i for i in commitments.check(conn, campaign_id=cid)["items"]
                    if "navy" in i["text"])
    assert ordinary["verdict"] == item["verdict"] == "unchecked"
    assert ordinary["what_it_means"] != item["what_it_means"]
    assert "count" in item["what_it_means"].lower()
    assert "count" not in ordinary["what_it_means"].lower()

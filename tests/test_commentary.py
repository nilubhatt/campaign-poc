"""
Product review defect 08 (P1): "Comments and speaker notes are never ingested."

The reviewer tested this against the PDFs already in the library and measured the loss:

    "One of them carries eight substantive /Text annotations — PowerPoint speaker notes
    exported as PDF comments. None of it is indexed. Querying the library with that second
    note almost verbatim returns the right record at only 0.63, matching on the summary a
    human wrote — not on the note. Had the note been ingested, a near-verbatim query would
    score above 0.9."

And said why it matters more here than the defect number suggests: "a partner deck returned
with tracked client comments is the feedback the library exists to remember, and it
currently has to be retyped by hand into the detail field to be captured at all."

The fix is explicit that commentary is a DIFFERENT LAYER, not more deck text: "store them as
a distinct layer from the deck body — not concatenated into deck_text. They are a different
kind of evidence: commentary carries author, date, and page or slide anchor, and it is
usually the internal reaction to the work rather than the work itself."
"""
import datetime
import json
import zipfile

import pytest
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.annotations import FreeText, Link, Text

import config
import core
import extract
import store


def _pdf_with_annotations(path, annots):
    """A PDF carrying real /Annots, built the way the reviewer's deck was: PowerPoint
    speaker notes exported as PDF comments."""
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=200, height=200)
    for page_index, annotation, author in annots:
        writer.add_annotation(page_number=page_index, annotation=annotation)
        # pypdf's helper does not write /T or /M, which is where author and date live.
        target = writer.pages[page_index]["/Annots"][-1].get_object()
        target[pypdf_name("/T")] = _pdf_string(author)
        target[pypdf_name("/M")] = _pdf_string("D:20260904143000+01'00'")
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def pypdf_name(value):
    from pypdf.generic import NameObject
    return NameObject(value)


def _pdf_string(value):
    from pypdf.generic import TextStringObject
    return TextStringObject(value)


NOTE_2 = ("Keep the leadership deck concise and put detail in the workbook. The workbook is "
          "the scalable input, and the deck is the summary of it.")
NOTE_3 = ("Call out the new competitor brand rule explicitly. These are compliance and brand "
          "safety rules, not preferences.")


# ── extraction ───────────────────────────────────────────────────────────────

def test_pdf_text_annotations_are_extracted_with_author_date_and_page(tmp_path):
    path = _pdf_with_annotations(tmp_path / "deck.pdf", [
        (1, Text(rect=(10, 10, 30, 30), text=NOTE_2), "Priya"),
        (2, Text(rect=(10, 10, 30, 30), text=NOTE_3), "Priya"),
    ])

    commentary, warnings = extract.extract_commentary(path)

    assert [c["text"] for c in commentary] == [NOTE_2, NOTE_3]
    assert [c["page"] for c in commentary] == [2, 3], "1-based, like every other anchor"
    assert all(c["author"] == "Priya" for c in commentary)
    assert commentary[0]["date"].startswith("2026-09-04"), "/M is a PDF date, not a string"
    assert commentary[0]["kind"] == "annotation"


def test_a_link_annotation_is_not_commentary(tmp_path):
    """"Every other deck in the library carries /Link annotations only, so no comment text
    is being silently dropped elsewhere" — a hyperlink is not somebody's opinion, and
    indexing one would put URL fragments into the evidence a judgment cites."""
    path = _pdf_with_annotations(tmp_path / "links.pdf", [
        (0, Link(rect=(10, 10, 30, 30), url="https://example.com/brief"), "Word"),
    ])

    commentary, _ = extract.extract_commentary(path)

    assert commentary == []


def test_a_freetext_annotation_is_commentary(tmp_path):
    path = _pdf_with_annotations(tmp_path / "ft.pdf", [
        (0, FreeText(text="Lose the third colourway", rect=(10, 10, 90, 30)), "Sam"),
    ])

    commentary, _ = extract.extract_commentary(path)

    assert commentary[0]["text"] == "Lose the third colourway"
    assert commentary[0]["author"] == "Sam"


def test_pptx_speaker_notes_are_extracted(tmp_path):
    prs = Presentation()
    for i in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = f"Speaker note {i + 1}: {NOTE_2}"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    commentary, _ = extract.extract_commentary(path)

    assert len(commentary) == 2
    assert commentary[0]["kind"] == "speaker_note"
    assert commentary[0]["slide"] == 1
    assert NOTE_2 in commentary[0]["text"]


def test_an_empty_notes_slide_is_not_a_comment(tmp_path):
    """PowerPoint creates a notesSlide for every slide the moment anything touches it, so
    the common case is an empty one. A library full of blank commentary rows is noise that
    dilutes every retrieval."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "   "
    path = tmp_path / "blank.pptx"
    prs.save(str(path))

    assert extract.extract_commentary(path)[0] == []


def test_pptx_reviewer_comments_are_extracted_with_author_and_timestamp(tmp_path):
    """The real prize: "a partner deck returned with tracked client comments is the feedback
    the library exists to remember". python-pptx has no API for these, so they are read out
    of the package parts directly — modernComments is what current PowerPoint writes."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "commented.pptx"
    prs.save(str(path))

    _add_modern_comment(path, author="Dana Ruiz", when="2026-09-04T14:30:00.000",
                        text="This is the adidas-affiliated profile we flagged in Q1.")

    commentary, _ = extract.extract_commentary(path)

    comment = next(c for c in commentary if c["kind"] == "comment")
    assert comment["author"] == "Dana Ruiz"
    assert comment["date"].startswith("2026-09-04")
    assert "adidas-affiliated" in comment["text"]


def test_a_deck_with_no_commentary_yields_none_without_complaining(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "plain.pptx"
    prs.save(str(path))

    commentary, warnings = extract.extract_commentary(path)

    assert commentary == []
    assert warnings == []


def test_unreadable_commentary_is_a_warning_not_a_failed_upload(tmp_path):
    """Commentary is the bonus layer. A malformed comments part must never cost the user the
    deck itself — they uploaded a brief, not a comments file."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "broken.pptx"
    prs.save(str(path))
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("ppt/modernComments/modernComment_1.xml", "<not-xml")

    commentary, warnings = extract.extract_commentary(path)

    assert commentary == []
    assert any(w["code"] == "commentary_unreadable" for w in warnings)


# ── it is a distinct layer, not more deck text ───────────────────────────────

def test_commentary_is_not_concatenated_into_the_deck_body(tmp_path, conn):
    """The review was explicit: "not concatenated into deck_text... usually the internal
    reaction to the work rather than the work itself". Merged in, a reviewer's objection
    reads back as something the brief itself claimed."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Colombia Q4 activation"
    slide.notes_slide.notes_text_frame.text = "I do not think the timeline is realistic."
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    result = core.ingest_campaign(conn, title="Colombia", asset_ref={"path": str(path)},
                                  confirm=True)
    campaign = store.get_campaign(conn, result["campaign_id"])

    assert "Colombia Q4 activation" in campaign["deck_text"]
    assert "not think the timeline" not in (campaign["deck_text"] or "")
    assert any("not think the timeline" in c["text"] for c in campaign["commentary"])


def test_commentary_carries_its_anchor_author_and_kind_into_storage(tmp_path, conn):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = NOTE_3
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    result = core.ingest_campaign(conn, title="Fabletics AI agent",
                                  asset_ref={"path": str(path)}, confirm=True)
    stored = store.get_campaign(conn, result["campaign_id"])["commentary"][0]

    assert stored["kind"] == "speaker_note"
    assert stored["anchor"] == "slide 1"
    assert stored["text"] == NOTE_3


def test_the_ingest_result_says_how_much_commentary_it_found(tmp_path, conn):
    """Silence is how defect 08 survived: nothing said the notes existed, so nothing said
    they were being dropped."""
    prs = Presentation()
    for _ in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = NOTE_2
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    result = core.ingest_campaign(conn, title="Noted", asset_ref={"path": str(path)},
                                  confirm=True)

    assert result["commentary_found"] == 3


# ── the reviewer's own measurement ───────────────────────────────────────────

def test_a_near_verbatim_query_on_a_note_finds_the_deck_that_carries_it(tmp_path, conn):
    """The review's numbers: querying near-verbatim on an unindexed note scored 0.63,
    matching the human-written summary instead. "Had the note been ingested, a near-verbatim
    query would score above 0.9."

    This pins the behaviour under the test embedder, which is not the same instrument the
    reviewer used — the claim was checked separately against the real one, end to end over
    MCP with Ollama, and returned 0.9929 (recorded in PRODUCT-REVIEW-PLAN.md §2.5). What
    this test is for is stopping the note going un-indexed again."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Fabletics AI agent — leadership update"
    slide.notes_slide.notes_text_frame.text = NOTE_2
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    core.ingest_campaign(conn, title="Fabletics AI agent", asset_ref={"path": str(path)},
                         confirm=True)

    # Near-verbatim, as the reviewer's probe was: the wording a person retypes from memory,
    # not a paraphrase. (A looser paraphrase of the same note lands around 0.90, which is
    # still the right record and still far above the 0.63 the un-indexed note managed.)
    hits = core.find_similar(conn, text="Keep the leadership deck concise and put detail in "
                                        "the workbook. The workbook is the scalable input "
                                        "and the deck summarises it.", top_k=3)

    assert hits, "the note is indexed, so it is findable"
    assert hits[0]["similarity"] > 0.9, f"scored {hits[0]['similarity']}"


def test_a_commentary_match_says_it_is_commentary_and_who_said_it(tmp_path, conn):
    """"tagged so retrieval can weigh or filter it" — a judgment that cites a reviewer's
    objection as if it were the brief's own claim is citing the wrong thing."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Colombia Q4"
    slide.notes_slide.notes_text_frame.text = ("The claw machine idea outperformed every "
                                               "other mechanic we tried in the UAE.")
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    core.ingest_campaign(conn, title="Colombia", asset_ref={"path": str(path)}, confirm=True)

    hit = core.find_similar(conn, text="claw machine mechanic performance", top_k=1)[0]

    assert hit["matched_kind"] == "commentary"
    assert hit["matched_anchor"] == "slide 1"


def test_commentary_can_be_excluded_from_retrieval(tmp_path, conn):
    """Weighing it means being able to not weigh it: "is this what the brief says, or what
    one person thought of it" is a question a marketer will ask."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Colombia Q4 influencer activation"
    slide.notes_slide.notes_text_frame.text = "A completely unrelated remark about logistics"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    core.ingest_campaign(conn, title="Colombia", asset_ref={"path": str(path)}, confirm=True)

    hit = core.find_similar(conn, text="unrelated remark about logistics", top_k=1,
                            include_commentary=False)

    assert hit and hit[0]["matched_kind"] == "body", (
        "the campaign still exists and still has a body chunk — excluding a layer narrows "
        "what can MATCH, it does not delete the record"
    )


def test_excluding_commentary_does_not_delete_the_campaign_from_the_results(tmp_path, conn):
    """Found by adversarial review, and it is the starvation bug the structured-filter path
    already carries a comment about — re-created one layer down. The filter ran AFTER the
    vector search's over-fetch window and after the per-campaign rollup, so a deck whose
    notes filled the window lost its body chunk before the filter ever saw it, and the
    campaign vanished from a search it plainly matches. The first version of the test above
    said `hit == [] or ...`, which blessed exactly this.

    The reviewer's own reproduction, kept verbatim in shape: the note text IS the query, so
    all 25 notes outrank the body chunk and fill the window ahead of it."""
    query = "unrelated remark about logistics"
    prs = Presentation()
    for i in range(25):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Logistics plan" if i == 0 else f"slide {i + 1}"
        slide.notes_slide.notes_text_frame.text = query
    path = tmp_path / "chatty.pptx"
    prs.save(str(path))

    core.ingest_campaign(conn, title="Logistics", deck_text=query,
                         asset_ref={"path": str(path)}, confirm=True)
    core.ingest_campaign(conn, title="Beach", deck_text="surf towels summer", confirm=True)

    assert core.find_similar(conn, text=query, top_k=1)[0]["title"] == "Logistics"

    body_only = core.find_similar(conn, text=query, top_k=1, include_commentary=False)

    assert body_only, "25 speaker notes filled the over-fetch window and the campaign vanished"
    assert body_only[0]["title"] == "Logistics"
    assert body_only[0]["matched_kind"] == "body"
    assert body_only[0]["similarity"] > 0.9


def test_the_filtered_path_never_had_the_window_problem(tmp_path, conn):
    """Control, and the proof that the defect was the over-fetch window rather than the
    filter itself: the structured-filter path ranks every candidate chunk, so the same query
    always found the body."""
    query = "unrelated remark about logistics"
    prs = Presentation()
    for i in range(25):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Logistics plan" if i == 0 else f"slide {i + 1}"
        slide.notes_slide.notes_text_frame.text = query
    path = tmp_path / "chatty.pptx"
    prs.save(str(path))
    core.ingest_campaign(conn, title="Logistics", deck_text=query,
                         asset_ref={"path": str(path)}, confirm=True)

    body_only = core.find_similar(conn, text=query, top_k=1, include_commentary=False,
                                  record_type="campaign")

    assert body_only and body_only[0]["matched_kind"] == "body"


# ── it must not cost the deck ────────────────────────────────────────────────

def test_a_database_from_before_this_release_gains_the_commentary_layer(tmp_path):
    """Every schema change from here on gets this test — 2.4 shipped one that was fine on a
    fresh database and impossible on a real one."""
    import sqlite3

    path = tmp_path / "old.db"
    old = sqlite3.connect(path)
    old.executescript("""
        CREATE TABLE campaigns (id TEXT PRIMARY KEY, title TEXT NOT NULL, collection TEXT,
                                markets TEXT NOT NULL DEFAULT '[]');
        CREATE TABLE campaign_chunks (id TEXT PRIMARY KEY, campaign_id TEXT NOT NULL,
                                      chunk_index INTEGER NOT NULL, text TEXT NOT NULL,
                                      embedded INTEGER NOT NULL DEFAULT 0,
                                      created_at REAL NOT NULL);
        CREATE TABLE evaluations (id TEXT PRIMARY KEY, campaign_id TEXT,
                                  subject_title TEXT NOT NULL, cited_ids TEXT,
                                  analysis TEXT NOT NULL, predictions TEXT,
                                  created_at REAL NOT NULL);
        CREATE TABLE reconciliations (id TEXT PRIMARY KEY, evaluation_id TEXT NOT NULL,
                                      actual TEXT, comparison TEXT NOT NULL,
                                      created_at REAL NOT NULL);
    """)
    old.execute("INSERT INTO campaign_chunks VALUES ('ch1','c1',0,'existing body text',1,1.0)")
    old.commit()
    old.close()

    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    store._migrate_schema(conn)

    row = conn.execute("SELECT * FROM campaign_chunks WHERE id='ch1'").fetchone()
    assert row["text"] == "existing body text"
    assert row["kind"] == "body", "an existing chunk is deck body, not commentary"


def test_a_deck_whose_commentary_cannot_be_read_still_ingests(tmp_path, conn, monkeypatch):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Colombia Q4"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    def boom(*a, **k):
        raise RuntimeError("comments part is corrupt")

    monkeypatch.setattr(extract, "extract_commentary", boom)

    result = core.ingest_campaign(conn, title="Colombia", asset_ref={"path": str(path)},
                                  confirm=True)

    assert result["chunks_total"] > 0
    assert any(w["code"] == "commentary_unreadable" for w in result["warnings"])


def test_commentary_is_capped_like_every_other_unbounded_input(tmp_path, monkeypatch):
    """A deck's comment count is controlled by whoever made the deck, and every other
    caller-sized loop in this server is bounded."""
    monkeypatch.setattr(config, "MAX_COMMENTARY_ITEMS", 5)
    prs = Presentation()
    for _ in range(config.MAX_COMMENTARY_ITEMS + 4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = NOTE_2
    path = tmp_path / "many.pptx"
    prs.save(str(path))

    commentary, warnings = extract.extract_commentary(path)

    assert len(commentary) == config.MAX_COMMENTARY_ITEMS
    assert any(str(config.MAX_COMMENTARY_ITEMS) in w["detail"] for w in warnings)


def _add_modern_comment(path, *, author, when, text):
    """Write a ppt/modernComments part into an existing .pptx, which is what PowerPoint
    does when a reviewer comments on a slide. python-pptx cannot author these."""
    import shutil

    src = str(path) + ".orig"
    shutil.move(str(path), src)
    author_id = "{6E9F5E3E-0000-0000-0000-000000000001}"
    comment_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p188:cmLst xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main">'
        f'<p188:cm id="1" authorId="{author_id}" created="{when}">'
        '<p188:txBody><a:bodyPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
        '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:r><a:t>{text}</a:t></a:r></a:p></p188:txBody></p188:cm></p188:cmLst>'
    )
    authors_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p188:cmAuthorLst xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main">'
        f'<p188:cmAuthor id="{author_id}" name="{author}" initials="DR"/></p188:cmAuthorLst>'
    )
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("ppt/modernComments/modernComment_1.xml", comment_xml)
        zout.writestr("ppt/authors.xml", authors_xml)


# ══ design review of 2.5 ═════════════════════════════════════════════════════

@pytest.fixture
def cited_records(conn):
    """§6.1 verifies a quote against the record it names, so these two have to exist and
    have to actually say what the findings below quote. They previously did not, which is
    the defect §6.1 closes."""
    import time
    now = time.time()
    conn.execute(
        "INSERT INTO campaigns (id, title, record_type, detail, created_at, updated_at) "
        "VALUES (?,?,?,?,?,?)",
        ("camp_peru", "Peru launch", "campaign",
         "Six-week flight across Lima and Arequipa.", now, now))
    conn.execute(
        "INSERT INTO campaigns (id, title, record_type, detail, created_at, updated_at) "
        "VALUES (?,?,?,?,?,?)",
        ("camp_x", "Jakarta launch", "campaign",
         "The flighting table gives a posting date per asset.", now, now))
    conn.commit()
    store.insert_chunks(conn, "camp_peru",
                        ["I do not think the timeline is realistic for a market this size."],
                        kind="commentary",
                        sources=[{"kind": "comment", "author": "Dana Ruiz",
                                  "anchor": "slide 4"}])


def test_a_finding_can_cite_a_comment_and_has_to_say_that_it_did(conn, cited_records):
    """The layer rule was stated on the browsing tool and absent on the judging one. A
    commentary chunk IS a retrieved chunk, so "I do not think the timeline is realistic" was
    a docstring-compliant quote for a finding against that campaign — and once saved it read
    permanently as something the campaign's own deck said. Worse, 6.1 will verify quotes
    against retrieved chunks, so without this the verification would bless the misattribution
    with a green tick."""
    result = core.save_evaluation(
        conn, subject_title="Colombia v2", verdict="revise",
        approve_if="The timeline is extended.",
        summary="Reuses a timeline the Peru team themselves flagged as unrealistic.",
        findings=[{
            "severity": "should_fix",
            "fix": "Change it",
            "kind": "precedent_departure",
            "departure": "regression",
            "finding": "Same six-week timeline Peru's own reviewer called unrealistic",
            "precedent": {"campaign_id": "camp_peru", "layer": "commentary",
                          "author": "Dana Ruiz", "anchor": "slide 4",
                          "quote": "I do not think the timeline is realistic"},
        }])

    cited = store.get_evaluation(conn, result["evaluation_id"])["findings"][0]["precedent"]
    assert cited["layer"] == "commentary"
    assert cited["author"] == "Dana Ruiz"
    assert cited["anchor"] == "slide 4"


def test_a_precedent_defaults_to_the_deck_body(conn, cited_records):
    """Unmarked means the brief itself, which is the safe reading: a citation that silently
    became commentary would be the defect this field exists to prevent."""
    result = core.save_evaluation(
        conn, subject_title="Colombia v2", verdict="revise", approve_if="It is fixed.", summary="Nothing is dated.",
        findings=[{"severity": "blocking", "fix": "Change it", "kind": "precedent_departure",
                   "departure": "regression", "finding": "No dates",
                   "precedent": {"campaign_id": "camp_x", "quote": "posting date per asset"}}])

    stored = store.get_evaluation(conn, result["evaluation_id"])
    assert stored["findings"][0]["precedent"]["layer"] == "body"


def test_an_unknown_layer_is_rejected(conn):
    with pytest.raises(ValueError) as exc:
        core.save_evaluation(
            conn, subject_title="T", verdict="revise", approve_if="It is fixed.", summary="s",
            findings=[{"severity": "blocking", "fix": "Change it", "kind": "precedent_departure",
                       "departure": "regression", "finding": "x",
                       "precedent": {"campaign_id": "c", "quote": "q", "layer": "hearsay"}}])

    assert "commentary" in str(exc.value)


def test_nothing_read_is_not_the_same_as_nothing_found(tmp_path, conn):
    """The images_checked lesson, re-learned: an empty list told a calling LLM nothing about
    whether anyone had looked, so it said "no reuse found" about slides it had never seen.
    commentary_found: 0 had exactly the same ambiguity on the deck_text-only path."""
    text_only = core.ingest_campaign(conn, title="Pasted", deck_text="A brief, as text.",
                                     confirm=True)

    assert text_only["commentary_found"] == 0
    assert text_only["commentary_checked"] is False, (
        "no file reached the server, so nothing was read — saying 'no comments' here would "
        "be a claim about a file nobody opened"
    )

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "plain.pptx"
    prs.save(str(path))
    with_file = core.ingest_campaign(conn, title="Real file",
                                     asset_ref={"path": str(path)}, confirm=True)

    assert with_file["commentary_found"] == 0
    assert with_file["commentary_checked"] is True, "read, and there genuinely were none"


def test_a_failed_read_does_not_report_itself_as_checked(tmp_path, conn, monkeypatch):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Colombia"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    monkeypatch.setattr(extract, "extract_commentary",
                        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("corrupt")))

    result = core.ingest_campaign(conn, title="Colombia", asset_ref={"path": str(path)},
                                  confirm=True)

    assert result["commentary_checked"] is False


def test_a_reply_is_its_own_comment_not_more_of_the_parents(tmp_path):
    """The scenario this item is named for — "a partner deck returned with tracked client
    comments" — is typically a client remark and an agency response. Joining every run of
    text under the comment node recorded the agency's answer as the client's words: defect
    08's own misattribution, one level down."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "thread.pptx"
    prs.save(str(path))
    _add_modern_comment_thread(path)

    commentary, _ = extract.extract_commentary(path)

    by_author = {c["author"]: c["text"] for c in commentary}
    assert by_author["Dana Ruiz"] == "Lose the third colourway."
    assert by_author["Sam Okafor"] == "Agreed, removing it."
    assert all("Agreed" not in t for a, t in by_author.items() if a == "Dana Ruiz")


def test_a_reply_says_what_it_is_replying_to(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "thread.pptx"
    prs.save(str(path))
    _add_modern_comment_thread(path)

    reply = next(c for c in extract.extract_commentary(path)[0] if c["author"] == "Sam Okafor")

    assert reply["reply_to"], "a reply with no parent is just an unattached remark"


def test_a_classic_pptx_comment_is_read_too(tmp_path):
    """ppt/comments/ is the older format and uses <p:text>, not the drawingml <a:t> that
    modernComments uses — so the docstring claimed both formats while one of them silently
    produced empty text and dropped the comment."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "classic.pptx"
    prs.save(str(path))
    _add_classic_comment(path, author="Priya Nair",
                         text="This is the roster we flagged in Q1.")

    comment = next(c for c in extract.extract_commentary(path)[0] if c["kind"] == "comment")

    assert comment["text"] == "This is the roster we flagged in Q1."
    assert comment["author"] == "Priya Nair"


def test_an_anchor_is_never_guessed_with_false_confidence(tmp_path):
    """A page index from a PDF is exact. A slide number scraped out of a comment part's
    FILENAME is a guess — for the classic format the number is the comment part's ordinal,
    not the slide's — and it was being presented in the same field, with the same confidence,
    as the exact one."""
    prs = Presentation()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "classic.pptx"
    prs.save(str(path))
    _add_classic_comment(path, author="Priya Nair", text="A remark.", part_index=1)

    comment = next(c for c in extract.extract_commentary(path)[0] if c["kind"] == "comment")

    assert comment["anchor"] == "deck" or comment["anchor"].startswith("slide"), comment
    if comment["anchor"].startswith("slide"):
        pytest.fail("the part filename is not the slide number; say 'deck' rather than "
                    "naming a slide the comment may not be on")


def test_commentary_can_be_narrowed_to_what_other_people_said(tmp_path, conn):
    """"What did the client say" is a different question from "what did anyone write
    anywhere", and a boolean could only ask the second."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Colombia Q4"
    slide.notes_slide.notes_text_frame.text = "Remember to mention the claw machine."
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    _add_modern_comment(path, author="Dana Ruiz", when="2026-09-04T14:30:00.000",
                        text="Remember the claw machine underperformed for us.")

    core.ingest_campaign(conn, title="Colombia", asset_ref={"path": str(path)}, confirm=True)

    hit = core.find_similar(conn, text="claw machine", top_k=1,
                            include_commentary=["comment"])[0]

    assert hit["matched_commentary_kind"] == "comment", "a speaker note is not a comment"
    assert hit["matched_author"] == "Dana Ruiz"


def test_the_page_or_slide_number_survives_as_a_number(tmp_path, conn):
    """The plan and the review both specify {page, author, date, text, kind}. Only the
    display string was kept, so nothing downstream could sort or group by position."""
    prs = Presentation()
    for i in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = f"Note on slide {i + 1}"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    result = core.ingest_campaign(conn, title="Noted", asset_ref={"path": str(path)},
                                  confirm=True)
    stored = store.get_campaign(conn, result["campaign_id"])["commentary"]

    assert [c["slide"] for c in stored] == [1, 2]


def _add_modern_comment_thread(path):
    """A client comment with an agency reply, the way PowerPoint nests them."""
    import shutil

    src = str(path) + ".orig"
    shutil.move(str(path), src)
    client = "{6E9F5E3E-0000-0000-0000-000000000001}"
    agency = "{6E9F5E3E-0000-0000-0000-000000000002}"
    ns = 'xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main"'
    a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

    def body(text):
        return (f'<p188:txBody><a:bodyPr {a}/><a:p {a}><a:r><a:t>{text}</a:t></a:r>'
                f'</a:p></p188:txBody>')

    comment_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p188:cmLst {ns}>'
        f'<p188:cm id="1" authorId="{client}" created="2026-09-04T14:30:00.000">'
        f'{body("Lose the third colourway.")}'
        f'<p188:replyLst>'
        f'<p188:reply id="2" authorId="{agency}" created="2026-09-04T16:05:00.000">'
        f'{body("Agreed, removing it.")}</p188:reply>'
        f'</p188:replyLst></p188:cm></p188:cmLst>'
    )
    authors_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p188:cmAuthorLst {ns}>'
        f'<p188:cmAuthor id="{client}" name="Dana Ruiz" initials="DR"/>'
        f'<p188:cmAuthor id="{agency}" name="Sam Okafor" initials="SO"/>'
        f'</p188:cmAuthorLst>'
    )
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("ppt/modernComments/modernComment_1.xml", comment_xml)
        zout.writestr("ppt/authors.xml", authors_xml)


def _add_classic_comment(path, *, author, text, part_index=1):
    """The pre-2018 format: ppt/comments/commentN.xml, <p:text>, author ids from
    ppt/commentAuthors.xml."""
    import shutil

    src = str(path) + ".orig"
    shutil.move(str(path), src)
    p = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    comment_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p:cmLst {p}><p:cm authorId="1" dt="2026-09-04T14:30:00" idx="1">'
        f'<p:pos x="100" y="100"/><p:text>{text}</p:text></p:cm></p:cmLst>'
    )
    authors_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p:cmAuthorLst {p}><p:cmAuthor id="1" name="{author}" initials="PN" '
        f'lastIdx="1" clrIdx="0"/></p:cmAuthorLst>'
    )
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(f"ppt/comments/comment{part_index}.xml", comment_xml)
        zout.writestr("ppt/commentAuthors.xml", authors_xml)


def test_a_deck_with_notes_but_no_body_text_still_keeps_the_notes(tmp_path, conn):
    """Adversarial review: the "nothing to embed" early return fires before the commentary
    is inserted, so a deck whose body extracted to nothing reported `commentary_found: 1`
    and stored zero chunks — the count described something that had just been thrown away."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = ("The four categories are valuable but "
                                               "incomplete without guardrails.")
    path = tmp_path / "notes_only.pptx"
    prs.save(str(path))

    result = core.ingest_campaign(conn, title="   ", asset_ref={"path": str(path)},
                                  confirm=True)

    assert result["commentary_found"] == 1
    stored = store.get_campaign(conn, result["campaign_id"])["commentary"]
    assert len(stored) == 1, "found and kept must mean the same thing"
    assert result["chunks_total"] == 1


def test_finish_indexing_completes_commentary_too(tmp_path, conn, monkeypatch):
    """Commentary chunks are embedded after the body ones, so they are the first thing a
    time budget drops — which makes them the thing most likely to need finishing."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Colombia Q4 activation plan for the southern region"
    for i in range(3):
        note = prs.slides.add_slide(prs.slide_layouts[6])
        note.notes_slide.notes_text_frame.text = f"Reviewer remark {i}: check the timeline"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.0)
    created = core.ingest_campaign(conn, title="Colombia", asset_ref={"path": str(path)},
                                   confirm=True)
    assert created["chunks_embedded"] < created["chunks_total"]

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 45.0)
    finished = core.finish_indexing(conn, campaign_id=created["campaign_id"])

    assert finished["complete"] is True
    hit = core.find_similar(conn, text="reviewer remark check the timeline", top_k=1,
                            include_commentary=["speaker_note"])
    assert hit and hit[0]["matched_kind"] == "commentary"


def test_a_malformed_annotation_date_is_kept_rather_than_invented(tmp_path):
    """`D:2026` became "2026--" and `D:99999999` became "9999-99-99" — a well-formed-looking
    date that no calendar contains. A date nobody can parse is still evidence of when
    somebody said something; a fabricated one is worse than none."""
    from pypdf.generic import TextStringObject

    path = _pdf_with_annotations(tmp_path / "dates.pdf", [
        (0, Text(rect=(10, 10, 30, 30), text="A remark"), "Priya"),
    ])
    import pypdf
    reader = pypdf.PdfReader(str(path))

    for raw, expected in (("D:2026", "D:2026"), ("D:99999999", "D:99999999"),
                          ("D:20260904143000+01'00'", "2026-09-04T14:30:00"),
                          ("D:20260904", "2026-09-04")):
        assert extract._pdf_date(TextStringObject(raw)) == expected, raw


def test_annotation_line_endings_are_readable(tmp_path):
    """PDF text uses a bare CR for a line break, which renders as one run-on line in any
    excerpt a marketer reads."""
    path = _pdf_with_annotations(tmp_path / "cr.pdf", [
        (0, Text(rect=(10, 10, 30, 30), text="line one\rline two"), "Priya"),
    ])

    assert extract.extract_commentary(path)[0][0]["text"] == "line one\nline two"

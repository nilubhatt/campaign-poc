"""
§9.3: check the named commitments, one by one.

The review: *"The richest signal is not statistical. A brief makes specific, checkable
promises: Colombia v2 commits to a claw machine, a photo booth, a matcha cart and a DJ; Peru
committed to thirty influencers in identical outfits; UAE commits to a 2000×1000mm matcha cart.
Fix: extract the commitment list at upload — it is mostly already in the deck as an experience
or floorplan list — and at delivery, check each one against the returned photos. 'Claw machine:
present. Photo booth: not visible in 14 delivered images. Single colourway: not observed, four
colourways present.' That is a post-mortem a marketer can act on, produced without anyone
writing it."*

**Read the review's own words for the negative case: "not visible in 14 delivered images".** Not
"absent", not "missing", not "was not delivered". That distinction is the whole honesty of this
item. The server can say what it could not see in the photographs it was given; it cannot say
the photo booth was not there. Fourteen photographs of a launch do not show everything at a
launch, and a product that reports "photo booth: absent" from their silence has asserted
something it has no way to know — about a supplier who may well have delivered it.

**Two different claims, two different bases.** That a LINE IS IN THE DECK is a fact — it is
quoted, and the quote is checkable. That the line is a COMMITMENT rather than a passing mention
is a reading, and that a photograph shows the thing is a resemblance score. So the commitment
carries the source line it was taken from, and the check is `heuristic`: §7.8's third basis,
which exists for exactly this — a threshold somebody chose.
"""
import pytest

import commitments
import core
import store

DECK = """Colombia v2 — launch experience

The activation floorplan includes:
- A claw machine loaded with branded merchandise
- A photo booth with instant prints
- A matcha cart serving from 11am
- A DJ from 6pm

Seeding: a single colourway across all thirty recipients.
"""


def _campaign(conn, title="Colombia v2", detail=DECK):
    return core.ingest_campaign(conn, title=title, market="LATAM", status="concluded",
                                detail=detail)["campaign_id"]


def _png(tmp_path, name, seed=1):
    import random

    from PIL import Image

    rng = random.Random(seed)
    img = Image.new("RGB", (64, 64), (250, 250, 250))
    px = img.load()
    for _ in range(700):
        x, y = rng.randrange(64), rng.randrange(64)
        c = (rng.randrange(256), rng.randrange(256), rng.randrange(256))
        for dx in range(4):
            for dy in range(4):
                if x + dx < 64 and y + dy < 64:
                    px[x + dx, y + dy] = c
    path = tmp_path / name
    img.save(path)
    return str(path)


# ── extraction: mostly already in the deck ──────────────────────────────────

def test_the_list_in_the_deck_becomes_the_commitment_list(conn):
    """"It is mostly already in the deck as an experience or floorplan list." A marketer
    should not retype what the brief already says."""
    cid = _campaign(conn)

    found = [c["text"] for c in commitments.for_campaign(conn, cid)]
    assert "A claw machine loaded with branded merchandise" in found
    assert "A DJ from 6pm" in found


def test_each_one_carries_the_line_it_came_from(conn):
    """That a line is IN THE DECK is a fact, and the quote is what makes it checkable. That
    the line is a commitment rather than a passing mention is a reading."""
    cid = _campaign(conn)

    claw = next(c for c in commitments.for_campaign(conn, cid) if "claw" in c["text"].lower())
    assert claw["source_line"] in DECK
    assert claw["basis"] == "computed", "the quote is; the reading is not"


def test_prose_that_is_not_a_list_is_not_a_commitment(conn):
    """A deck is mostly prose. Reading every sentence as a promise is the "forty new keys"
    shape: a list that fires on everything is one nobody works through."""
    cid = _campaign(conn, "Plain", detail="A launch in Colombia. It went well. We were "
                                          "pleased with the turnout and the press coverage.")

    assert commitments.for_campaign(conn, cid) == []


def test_the_number_of_commitments_is_bounded(conn):
    """A floorplan list is a handful of things. A deck with two hundred bullets is a deck, not
    a commitment list, and reading it as one produces a post-mortem nobody finishes."""
    cid = _campaign(conn, "Long", detail="Plan:\n" + "\n".join(
        f"- Item number {n} in the plan" for n in range(200)))

    assert len(commitments.for_campaign(conn, cid)) <= commitments.MAX_COMMITMENTS


def test_a_commitment_can_be_added_by_hand(conn):
    """"Thirty influencers in identical outfits" is a promise in a sentence, not a bullet."""
    cid = _campaign(conn, "Peru", detail="A seeding campaign.")

    added = commitments.add(conn, campaign_id=cid,
                            text="Thirty influencers in identical outfits",
                            source_line="Seeded to thirty influencers, identical outfits.")

    assert [c["text"] for c in commitments.for_campaign(conn, cid)] == [added["text"]]


def test_one_that_is_not_a_commitment_can_be_dropped(conn):
    """Mechanical extraction will pick up lines that are not promises. A list nobody can
    correct is a list that stops being read."""
    cid = _campaign(conn)
    first = commitments.for_campaign(conn, cid)[0]

    commitments.drop(conn, first["id"], why="Not a deliverable.")

    assert first["id"] not in [c["id"] for c in commitments.for_campaign(conn, cid)]
    assert commitments.describe(conn, first["id"])["status"] == "dropped", "kept, not deleted"


# ── the check, and what it may claim ────────────────────────────────────────

def test_a_commitment_the_photographs_show_is_reported_present(conn, tmp_path):
    cid = _campaign(conn)
    shot = core.ingest_image_asset(conn, campaign_id=cid,
                                   asset_ref={"path": _png(tmp_path, "claw.png")},
                                   phase="delivered", captured_on="2026-03-14")
    _stage_all(conn, cid)
    claw = next(c for c in commitments.for_campaign(conn, cid) if "claw" in c["text"].lower())
    _make_visible(conn, claw["id"], shot["asset_id"])

    report = commitments.check(conn, campaign_id=cid)
    item = next(i for i in report["items"] if i["commitment_id"] == claw["id"])

    assert item["verdict"] == "present"
    assert item["seen_in"]["asset_id"] == shot["asset_id"]


def test_a_commitment_nothing_shows_is_never_called_absent(conn, tmp_path):
    """The review's own words: "Photo booth: NOT VISIBLE IN 14 DELIVERED IMAGES." Fourteen
    photographs of a launch do not show everything at a launch, and a product reporting
    "absent" from their silence has asserted something it cannot know — about a supplier who
    may well have delivered it."""
    cid = _campaign(conn)
    for n in range(3):
        core.ingest_image_asset(conn, campaign_id=cid,
                                asset_ref={"path": _png(tmp_path, f"s{n}.png", seed=n + 2)},
                                phase="delivered", captured_on="2026-03-14")

    _stage_all(conn, cid)
    report = commitments.check(conn, campaign_id=cid)
    booth = next(i for i in report["items"] if "booth" in i["text"].lower())

    assert booth["verdict"] == "not_visible"
    assert "not visible in 3 delivered images" in booth["what_it_means"]
    for word in ("absent", "missing", "was not delivered", "did not happen"):
        assert word not in booth["what_it_means"].lower()


def test_the_check_is_a_resemblance_and_says_so(conn, tmp_path):
    """A photograph showing a claw machine is a resemblance score, not a fact. §7.8's third
    basis exists for exactly this."""
    cid = _campaign(conn)
    core.ingest_image_asset(conn, campaign_id=cid,
                            asset_ref={"path": _png(tmp_path, "a.png")},
                            phase="delivered", captured_on="2026-03-14")

    report = commitments.check(conn, campaign_id=cid)
    assert report["basis"] == "heuristic"
    assert all(item["basis"] == "heuristic" for item in report["items"])


def test_nothing_delivered_is_not_everything_missing(conn):
    """A campaign whose photographs have not come back has not failed its commitments. This is
    §9.2's rule in the other half of the phase: an empty set is not a result."""
    cid = _campaign(conn)

    report = commitments.check(conn, campaign_id=cid)

    assert report["status"] == "nothing_to_check"
    assert all(item["verdict"] == "unchecked" for item in report["items"])
    assert "has come back" in report["what_it_means"]


def test_a_campaign_with_no_commitments_says_so(conn, tmp_path):
    cid = _campaign(conn, "Plain", detail="A launch in Colombia that went well.")
    core.ingest_image_asset(conn, campaign_id=cid,
                            asset_ref={"path": _png(tmp_path, "a.png")},
                            phase="delivered", captured_on="2026-03-14")

    report = commitments.check(conn, campaign_id=cid)
    assert report["items"] == []
    assert "nothing the brief named" in report["what_it_means"]


def test_it_reads_as_the_post_mortem_the_review_describes(conn, tmp_path):
    """"Claw machine: present. Photo booth: not visible in 14 delivered images." One line per
    promise, in the brief's own words."""
    cid = _campaign(conn)
    shot = core.ingest_image_asset(conn, campaign_id=cid,
                                   asset_ref={"path": _png(tmp_path, "claw.png")},
                                   phase="delivered", captured_on="2026-03-14")
    _stage_all(conn, cid)
    claw = next(c for c in commitments.for_campaign(conn, cid) if "claw" in c["text"].lower())
    _make_visible(conn, claw["id"], shot["asset_id"])

    lines = commitments.check(conn, campaign_id=cid)["summary"]
    assert any(line.startswith("A claw machine") and "present" in line for line in lines)
    assert any("not visible" in line for line in lines)


def test_the_vision_model_being_off_is_not_a_verdict(conn, tmp_path):
    """§6.4's lesson, again: an outage that reads as "none of your commitments appeared" is
    the worst possible false negative — it accuses a supplier."""
    cid = _campaign(conn)
    core.ingest_image_asset(conn, campaign_id=cid,
                            asset_ref={"path": _png(tmp_path, "a.png")},
                            phase="delivered", captured_on="2026-03-14")
    conn.execute("DELETE FROM asset_vectors")
    conn.commit()

    report = commitments.check(conn, campaign_id=cid)
    assert all(item["verdict"] == "unchecked" for item in report["items"])
    assert report["warnings"]


def test_it_reaches_the_model_over_the_protocol(conn, tmp_path):
    import asyncio
    import json

    import mcp_server

    async def call(name, args):
        return json.loads((await mcp_server.mcp.call_tool(name, args)).content[0].text)

    cid = _campaign(conn)
    core.ingest_image_asset(conn, campaign_id=cid,
                            asset_ref={"path": _png(tmp_path, "a.png")},
                            phase="delivered", captured_on="2026-03-14")

    report = asyncio.run(call("check_commitments", {"campaign_id": cid}))

    assert report["basis"] == "heuristic"
    assert len(report["items"]) == 4


def _stage_all(conn, campaign_id):
    """Give every commitment a real, distinct vector.

    The `hash` provider returns zeros from `embed_text` — honestly, because it has no shared
    text/image space — so without this the verdicts are all `unchecked` and the REPORTING is
    unreachable. Staging is how these tests exercise the reporting without pretending the hash
    provider can recognise a claw machine (D124).
    """
    import random

    import config
    import vectorstore

    vectorstore.init(conn, space="commitment")
    for n, promise in enumerate(commitments.for_campaign(conn, campaign_id)):
        rng = random.Random(1000 + n)
        vec = [rng.uniform(-1, 1) for _ in range(config.CLIP_EMBED_DIM)]
        vectorstore.add(conn, f"commitment:{promise['id']}", vec, space="commitment")


def _make_visible(conn, commitment_id, asset_id):
    """Put the commitment's text vector and the asset's image vector in the same place.

    The `hash` CLIP provider the suite runs on has no shared text/image space — a hash of a
    string and a hash of pixels are not comparable, whatever the numbers do. So visibility is
    staged directly rather than pretended: what these tests check is the REPORTING, and
    whether real CLIP recognises a claw machine is a question for a fixture of recorded
    vectors, which is D124.
    """
    import vectorstore

    vec = vectorstore.get_many(conn, [asset_id], space="asset")[asset_id]
    vectorstore.init(conn, space="commitment")
    vectorstore.add(conn, f"commitment:{commitment_id}", vec, space="commitment")


def test_the_photograph_that_makes_the_check_possible_offers_it(conn, tmp_path):
    """D116, the fifth time. The review calls this "the richest signal", and it was reachable
    from nothing but its own definition."""
    cid = _campaign(conn)

    landed = core.ingest_image_asset(conn, campaign_id=cid,
                                     asset_ref={"path": _png(tmp_path, "a.png")},
                                     phase="delivered", captured_on="2026-03-14")

    assert landed["next_actions"][0]["tool"] == "check_commitments"
    assert landed["next_actions"][0]["prefilled_args"]["campaign_id"] == cid


def test_a_brief_that_promised_nothing_specific_offers_no_check(conn, tmp_path):
    cid = _campaign(conn, "Plain", detail="A launch that went well.")
    landed = core.ingest_image_asset(conn, campaign_id=cid,
                                     asset_ref={"path": _png(tmp_path, "a.png")},
                                     phase="delivered", captured_on="2026-03-14")

    assert [o["tool"] for o in landed["next_actions"]] == []


def test_a_commitment_added_without_a_source_is_refused(conn):
    """A commitment that cannot show where it came from is one nobody can check — and a
    post-mortem resting on it can say only "the library says so"."""
    cid = _campaign(conn, "Peru", detail="A seeding campaign.")

    with pytest.raises(ValueError) as e:
        commitments.add(conn, campaign_id=cid, text="Thirty influencers", source_line="  ")
    assert "source_line" in str(e.value)


def test_nothing_delivered_is_unchecked_even_when_the_phrase_embeds(conn):
    """The staged vector is what makes this test mean anything: with the `hash` provider
    `embed_text` returns zeros, so every verdict is `unchecked` for the wrong reason and the
    guard that produces the right one is unreachable."""
    cid = _campaign(conn)
    _stage_all(conn, cid)

    report = commitments.check(conn, campaign_id=cid)

    assert all(i["verdict"] == "unchecked" for i in report["items"])
    assert "has come back" in report["what_it_means"]
    assert all("not visible" not in i["what_it_means"] for i in report["items"])


def test_an_unindexed_photograph_is_unchecked_even_when_the_phrase_embeds(conn, tmp_path):
    """The worst possible false negative in this product: an outage that reads as "none of your
    commitments appeared" ACCUSES A SUPPLIER. Same masking as above — without a staged vector
    this passes whether the guard is there or not."""
    cid = _campaign(conn)
    core.ingest_image_asset(conn, campaign_id=cid,
                            asset_ref={"path": _png(tmp_path, "a.png")},
                            phase="delivered", captured_on="2026-03-14")
    _stage_all(conn, cid)
    conn.execute("DELETE FROM asset_vectors")
    conn.commit()

    report = commitments.check(conn, campaign_id=cid)

    assert all(i["verdict"] == "unchecked" for i in report["items"])
    assert report["counts"]["not_visible"] == 0
    assert report["warnings"]


def test_brand_guidelines_do_not_promise_anything(conn):
    """A reference record is full of bulleted lists and none of them is a promise THIS campaign
    made. §9.2 learned the same thing about which RECORD a check runs on, one item ago."""
    ref = core.ingest_campaign(conn, title="Brand guidelines", record_type="reference",
                               market="LATAM", detail=DECK)["campaign_id"]

    assert commitments.for_campaign(conn, ref) == []


def test_an_edited_brief_promises_what_it_now_says(conn):
    """Checking a new deck against the old deck's promises is "learning from the wrong
    document" — the failure this whole phase exists to stop, arriving inside the fix for it."""
    cid = _campaign(conn)
    assert any("claw" in c["text"].lower() for c in commitments.for_campaign(conn, cid))

    core.update_campaign(conn, campaign_id=cid,
                         detail="Plan:\n- A silent disco\n- A flower wall\n")

    texts = [c["text"] for c in commitments.for_campaign(conn, cid)]
    assert texts == ["A silent disco", "A flower wall"]


def test_an_edit_does_not_undo_what_a_person_decided(conn):
    """A commitment somebody added by hand stays; one somebody dropped stays dropped rather
    than reappearing on the next edit."""
    cid = _campaign(conn)
    booth = next(c for c in commitments.for_campaign(conn, cid) if "booth" in c["text"].lower())
    commitments.drop(conn, booth["id"], why="Cut from the plan.")
    commitments.add(conn, campaign_id=cid, text="Thirty influencers",
                    source_line="Seeded to thirty influencers.")

    core.update_campaign(conn, campaign_id=cid, detail=DECK)

    texts = [c["text"] for c in commitments.for_campaign(conn, cid)]
    assert "Thirty influencers" in texts, "what a person added survives"
    assert not any("booth" in t.lower() for t in texts), "and what they dropped stays dropped"


def test_a_deleted_campaign_leaves_no_commitment_vectors_behind(conn, tmp_path):
    import vectorstore

    cid = _campaign(conn)
    _stage_all(conn, cid)
    one = commitments.for_campaign(conn, cid)[0]["id"]
    assert vectorstore.get_many(conn, [f"commitment:{one}"], space="commitment")

    store.delete_campaign(conn, cid)

    assert vectorstore.get_many(conn, [f"commitment:{one}"], space="commitment") == {}


# ── what two reviews found ──────────────────────────────────────────────────

def test_a_real_deck_yields_its_list(conn, tmp_path):
    """The feature returned NOTHING on the one path §9.3 was written for. In OOXML a bullet is
    paragraph formatting, and usually inherited from the layout rather than written on the
    paragraph at all — so `text_frame.text` hands back the words with no list-ness, and the
    deck's own experience list produced zero commitments."""
    import extract
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Colombia v2 — Experience"
    body = slide.placeholders[1].text_frame
    body.text = "A claw machine loaded with branded merchandise"
    for line in ("A photo booth with instant prints", "A DJ from 6pm"):
        body.add_paragraph().text = line
    deck = tmp_path / "deck.pptx"
    prs.save(deck)

    units, _warnings = extract.extract_units(deck)
    found, _skipped = commitments._candidate_lines("\n".join(units))

    assert [f["text"] for f in found] == [
        "A claw machine loaded with branded merchandise",
        "A photo booth with instant prints", "A DJ from 6pm"]


def test_one_line_of_body_text_is_not_a_list(conn, tmp_path):
    """The marker goes into the text this library stores, quotes against and searches. A body
    placeholder inherits a bullet whatever is in it, so a single paragraph would be marked as a
    list item — one line under a heading is a statement."""
    import extract
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Colombia"
    slide.placeholders[1].text_frame.text = "A launch that went well."
    deck = tmp_path / "one.pptx"
    prs.save(deck)

    units, _ = extract.extract_units(deck)
    assert "- A launch" not in units[0]


def test_an_agenda_is_not_a_list_of_promises(conn):
    """"Welcome and introductions: not visible in 14 delivered images" is a post-mortem line
    about a meeting, and "Maria Gonzalez, Account Director: not visible" is one about a
    person. Agenda and team slides are slides 1 and 2 of nearly every deck."""
    cid = _campaign(conn, "Agenda deck", detail=(
        "Agenda\n- Welcome and introductions\n- The work\n\n"
        "The team\n- Maria Gonzalez, Account Director\n\n"
        "Timeline\n- Jan 12 — kick-off\n\n"
        "Activation floorplan\n- A claw machine loaded with branded merchandise\n"))

    assert [c["text"] for c in commitments.for_campaign(conn, cid)] == [
        "A claw machine loaded with branded merchandise"]


def test_a_promise_about_how_many_is_never_answered_by_a_resemblance(conn, tmp_path):
    """The review's own third example, and the most damaging thing this module could produce.
    "Single colourway across all recipients" run through "which photograph most resembles this
    phrase" is guaranteed to CONFIRM itself: more colourways means more hero images means a
    higher maximum. The verdict would read as the promise being kept, on exactly the evidence
    that it was broken."""
    cid = _campaign(conn, "Peru", detail=(
        "Seeding\n- Single colourway across all recipients\n- A branded tote\n"))
    core.ingest_image_asset(conn, campaign_id=cid,
                            asset_ref={"path": _png(tmp_path, "a.png")},
                            phase="delivered", captured_on="2026-03-14")
    _stage_all(conn, cid)

    report = commitments.check(conn, campaign_id=cid)
    single = next(i for i in report["items"] if "colourway" in i["text"].lower())

    assert single["verdict"] == "unchecked"
    assert "how MANY or about ALL" in single["what_it_means"]
    assert "Somebody has to look" in single["what_it_means"]


def test_a_photograph_has_to_beat_the_phrases_own_median(conn, tmp_path):
    """A single absolute cutoff cannot exist: CLIP constrains only RELATIVE order, so every
    phrase has its own floor — measured against the real weights, "Jan 12 — kick-off" scored
    0.246 against a solid green square while a genuine promise scored 0.231 against noise."""
    assert commitments.MARGIN > 0
    cid = _campaign(conn)
    for n in range(4):
        core.ingest_image_asset(conn, campaign_id=cid,
                                asset_ref={"path": _png(tmp_path, f"s{n}.png", seed=n + 5)},
                                phase="delivered", captured_on="2026-03-14")
    _stage_all(conn, cid)
    # A phrase that scores the same against everything is not describing any of them.
    import config
    import vectorstore
    flat = [0.5] * config.CLIP_EMBED_DIM
    for asset in store.assets_in_phase(conn, cid, "delivered"):
        conn.execute("DELETE FROM asset_vectors WHERE vector_id = ?", (asset["id"],))
        vectorstore.add(conn, asset["id"], flat, space="asset")
    claw = next(c for c in commitments.for_campaign(conn, cid) if "claw" in c["text"].lower())
    vectorstore.add(conn, f"commitment:{claw['id']}", flat, space="commitment")
    conn.commit()

    item = next(i for i in commitments.check(conn, campaign_id=cid)["items"]
                if i["commitment_id"] == claw["id"])
    assert item["similarity"] > commitments.VISIBLE_AT, "the absolute test alone would pass it"
    assert item["verdict"] == "not_visible", "but it resembles every photograph equally"


def test_every_verdict_carries_the_closest_photographs(conn, tmp_path):
    """"Here are the three photos most like each promise; you look" is true whatever the
    threshold does — and computing the ranking and keeping only the maximum is the same waste
    §9.2 was pulled up for one item ago."""
    cid = _campaign(conn)
    for n in range(4):
        core.ingest_image_asset(conn, campaign_id=cid,
                                asset_ref={"path": _png(tmp_path, f"s{n}.png", seed=n + 5)},
                                phase="delivered", captured_on="2026-03-14")
    _stage_all(conn, cid)

    item = commitments.check(conn, campaign_id=cid)["items"][0]
    assert len(item["closest"]) == 3
    assert all(c["similarity"] is not None for c in item["closest"])
    assert "if you want to look" in item["what_it_means"]


def test_nothing_checked_never_reads_as_nothing_found(conn, tmp_path):
    """The supplier-accusing false negative, produced by the HEADLINE rather than by any
    verdict: every item came back `unchecked` and the report still said "0 of 4 named promises
    can be seen... the rest were not visible in them"."""
    cid = _campaign(conn, "Peru", detail="Seeding\n- Single colourway across all recipients\n")
    core.ingest_image_asset(conn, campaign_id=cid,
                            asset_ref={"path": _png(tmp_path, "a.png")},
                            phase="delivered", captured_on="2026-03-14")

    report = commitments.check(conn, campaign_id=cid)

    assert report["status"] == "nothing_to_check"
    assert "not visible in them" not in report["what_it_means"]
    assert "Nothing is known about any of them" in report["what_it_means"]


def test_the_summary_says_how_many_were_actually_searched(conn, tmp_path):
    """With one of three photographs unindexed the summary said "not visible in 3 delivered
    images" while the item said 2 — and the summary is the line the review calls actionable."""
    cid = _campaign(conn)
    for n in range(3):
        core.ingest_image_asset(conn, campaign_id=cid,
                                asset_ref={"path": _png(tmp_path, f"s{n}.png", seed=n + 5)},
                                phase="delivered", captured_on="2026-03-14")
    _stage_all(conn, cid)
    one = store.assets_in_phase(conn, cid, "delivered")[0]["id"]
    conn.execute("DELETE FROM asset_vectors WHERE vector_id = ?", (one,))
    conn.commit()

    report = commitments.check(conn, campaign_id=cid)
    assert report["searched"] == 2
    assert all("in 2 delivered images" in line for line in report["summary"]
               if "not visible" in line)


def test_a_title_typo_does_not_churn_every_commitment_id(conn):
    """Re-extraction deletes the open extracted rows and reinserts them with new ids, so every
    `commitment_id` in a list the model had just shown went stale and `drop_commitment` failed
    with "not a commitment on file". The promises are in the body, not the title."""
    cid = _campaign(conn)
    before = [c["id"] for c in commitments.for_campaign(conn, cid)]

    core.update_campaign(conn, campaign_id=cid, title="Colombia v2 (final)")

    assert [c["id"] for c in commitments.for_campaign(conn, cid)] == before


def test_the_list_is_shown_at_upload_before_it_accuses_anybody(conn):
    """D116's sixth occurrence, in the item after the one that named the process gap. A
    mechanically-extracted list first reached a human as verdicts in a post-mortem, after the
    junk entries had already said something about a supplier."""
    result = core.ingest_campaign(conn, title="Colombia v2", market="LATAM",
                                  status="concluded", detail=DECK)

    shown = result["commitments"]
    assert shown["count"] == 4
    assert shown["promises"][0]["source_line"]
    assert {a["tool"] for a in shown["next_actions"]} == {"list_commitments", "add_commitment"}


def test_a_check_with_nothing_named_offers_the_way_to_fix_it(conn, tmp_path):
    cid = _campaign(conn, "Plain", detail="A launch that went well.")
    core.ingest_image_asset(conn, campaign_id=cid,
                            asset_ref={"path": _png(tmp_path, "a.png")},
                            phase="delivered", captured_on="2026-03-14")

    report = commitments.check(conn, campaign_id=cid)
    assert report["next_actions"][0]["tool"] == "add_commitment"


def test_a_wrong_path_says_so_instead_of_crashing(conn):
    """Every one of these was a bare string in a warnings list that `notices.collapse` reads as
    dicts — so an ordinary wrong path crashed `upload_campaign` with `TypeError`, which is not
    a ValueError and reached the model as "Error executing tool" with the reason discarded."""
    out = core.ingest_campaign(conn, title="X", detail="d",
                               asset_ref={"path": "/nope/missing.pdf"})

    assert out["campaign_id"], "the record is still stored"
    assert any(w["code"] == "asset_unreadable" for w in out["warnings"])
    assert "/nope/missing.pdf" in out["warnings"][0]["detail"]

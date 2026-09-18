"""
§13.4 — the three things the re-index left behind.

    D99   `embedding.embed` on the re-index path is called with NO timeout, unlike ingest
          (`timeout=remaining`) and `finish_indexing` (`timeout=remaining_time`), so a hung
          embedder hangs the handler once per chunk. The time-budget machinery already exists;
          this call site is the one that does not use it.
    D100  Chunk indices are allocated `MAX(chunk_index) + 1` across the WHOLE campaign, so a
          body that gains a position lands after the commentary. Measured before the change:
          a three-section deck with one comment, grown to six sections, came back as
          `body 0-6, commentary 7, body 8-13` — the body's own sequence interrupted by a
          different layer. Filed as latent because nothing read `chunk_index == 0`; §13.3 made
          body order load-bearing for incremental rebuild, so it is not latent any more.
    D109  `corrections._looks_like` is lexical — stemmed content words, overlap against the
          shorter rule — so it misses a paraphrase sharing no vocabulary. Deferred BECAUSE of
          D99: `note` is a write path, and adding a model call to a write before the embedder
          has a timeout trades a missed suggestion for a hung tool. That is the dependency
          this item exists to unblock, and it is unblocked by fixing D99 first.

The invariant §13.3 handed this item, and which it must not break: `_rebuild_body_index`
matches `store.body_chunks(... ORDER BY chunk_index)` positionally against `chunking.pack`'s
output. Any renumbering that does not preserve the body's RELATIVE order makes every position
mismatch, and the failure is silent — a total mismatch just re-embeds the whole deck, which is
the old behaviour.
"""
import pytest

import core
import corrections
import embedding
import store


_PARA = "The Bogota launch runs paid social and out-of-home this quarter. " * 30


def _deck(sections):
    return "\n\n".join(f"Section {i}. {_PARA}" for i in range(sections))


def _record(conn, sections=3):
    return core.ingest_campaign(conn, title="Bogota launch", market="Peru",
                                status="concluded", deck_text=_deck(sections),
                                confirm=True)["campaign_id"]


def _indices(conn, cid):
    return [(r["chunk_index"], r["kind"]) for r in conn.execute(
        "SELECT chunk_index, kind FROM campaign_chunks WHERE campaign_id = ? "
        "ORDER BY chunk_index", (cid,))]


# ---------------------------------------------------------------------------------------
# D99 — the one embed path with no deadline


def test_the_re_index_embeds_against_a_deadline(conn, monkeypatch):
    """Ingest passes `timeout=remaining` and `finish_indexing` passes `timeout=remaining_time`.
    The re-index passed nothing, so an embedder that hangs hangs the handler once per chunk —
    and a re-index is the path that runs over every chunk of a deck."""
    cid = _record(conn, sections=3)
    conn.execute("UPDATE campaigns SET deck_text = ? WHERE id = ?", (_deck(5), cid))
    conn.commit()

    given = []
    real = embedding.embed
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (given.append(timeout),
                                                    real(text, timeout=timeout))[1])

    core._rebuild_body_index(conn, cid)

    assert given, "the fixture did not reach the embedder"
    assert all(t is not None for t in given), (
        "the re-index embeds with no timeout, so a hung embedder hangs the handler per chunk"
    )


def test_a_slow_re_index_stops_at_the_budget_rather_than_running_on(conn, monkeypatch):
    """The budget is a ceiling on the HANDLER, not on each call — §2.1's whole point about
    partial state. What is left undone is reported and `finish_indexing` closes it."""
    import time as _time

    cid = _record(conn, sections=6)
    # Disjoint vocabulary, so "did this chunk get rewritten" is decidable — `chunking.pack`
    # splits a long paragraph and the continuation carries no section header.
    fresh = "\n\n".join(f"Chapter {i}. " + ("wholly different wording here. " * 30)
                        for i in range(6))
    conn.execute("UPDATE campaigns SET deck_text = ? WHERE id = ?", (fresh, cid))
    conn.commit()

    import config
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.2)
    real = embedding.embed
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (_time.sleep(0.08),
                                                    real(text, timeout=timeout))[1])

    started = _time.perf_counter()
    out = core._rebuild_body_index(conn, cid)
    elapsed = _time.perf_counter() - started

    assert elapsed < 2.0, f"a re-index ran for {elapsed:.1f}s against a 0.2s budget"
    assert out["embedded"] < out["chunks"], "the fixture has to leave something undone"
    # And the WORDS are all right even though the index is behind — §2.1's partial state,
    # and the property §13.3's two-pass rebuild exists to guarantee.
    assert not any("Bogota launch runs paid social" in r["text"]
                   for r in store.body_chunks(conn, cid)), (
        "chunks still hold the old deck after the budget ran out"
    )


def test_one_tool_call_spends_one_budget(conn, tmp_path, monkeypatch):
    """`attach_deck` rebuilds the body, embeds the commentary and indexes the images. Three
    independent budgets meant one tool call could spend three times the number the config
    names — and `ingest_campaign` already threads a SINGLE deadline through exactly those
    three phases, so the question had been settled and this had quietly answered it
    differently. Both reviewers measured it."""
    import time as _time

    from pptx import Presentation

    import config

    cid = core.ingest_campaign(conn, title="Bogota launch", market="Peru", status="concluded",
                               detail="Typed from memory.", confirm=True)["campaign_id"]
    deck = tmp_path / "bogota.pptx"
    show = Presentation()
    for n in range(6):
        slide = show.slides.add_slide(show.slide_layouts[1])
        slide.shapes.title.text = f"Section {n}"
        slide.placeholders[1].text = _PARA
    show.save(deck)

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.3)
    real = embedding.embed
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (_time.sleep(0.06),
                                                    real(text, timeout=timeout))[1])

    started = _time.perf_counter()
    core.attach_deck(conn, campaign_id=cid, asset_ref={"path": str(deck)})
    elapsed = _time.perf_counter() - started

    assert elapsed < 3 * config.TOOL_TIME_BUDGET_SECONDS, (
        f"one tool call took {elapsed:.2f}s against a {config.TOOL_TIME_BUDGET_SECONDS}s budget"
    )


# ---------------------------------------------------------------------------------------
# D100 — a layer's indices interrupted by another layer


def test_a_growing_deck_keeps_its_body_chunks_together(conn):
    """D100. Indices were allocated `MAX + 1` over the WHOLE campaign, so a body that gains a
    position landed after the commentary: `body 0-6, commentary 7, body 8-13`. One layer's
    sequence interrupted by another's."""
    cid = _record(conn, sections=3)
    store.insert_chunks(conn, cid, ["A client remark."], kind="commentary",
                        sources=[{"kind": "comment", "author": "R. Vega"}])
    conn.execute("UPDATE campaigns SET deck_text = ? WHERE id = ?", (_deck(6), cid))
    conn.commit()

    core._rebuild_body_index(conn, cid)

    body = [i for i, kind in _indices(conn, cid) if kind == "body"]
    assert body == list(range(len(body))), (
        f"the body's indices are {body} — not a contiguous run from 0"
    )


def test_chunk_zero_is_the_summary_even_when_commentary_came_first(conn):
    """"Chunk 0 is the summary" is the invariant D100 says was silently lost. A record that
    gained its commentary BEFORE its body — which is what attaching a deck to a record
    somebody had already annotated looks like — put a client's remark at index 0."""
    cid = core.ingest_campaign(conn, title="Bogota launch", market="Peru", status="concluded",
                               detail="Typed from memory.", confirm=True)["campaign_id"]
    store.insert_chunks(conn, cid, ["A client remark."], kind="commentary",
                        sources=[{"kind": "comment", "author": "R. Vega"}])
    conn.execute("UPDATE campaigns SET deck_text = ? WHERE id = ?", (_deck(3), cid))
    conn.commit()

    core._rebuild_body_index(conn, cid)

    first = store.body_chunks(conn, cid)[0]
    assert first["chunk_index"] == 0, f"the summary sits at index {first['chunk_index']}"
    assert "Bogota launch" in first["text"], "chunk 0 is not the title-and-detail summary"


def test_renumbering_does_not_break_the_incremental_rebuild(conn, monkeypatch):
    """The invariant §13.3 handed this item. `_rebuild_body_index` matches stored chunks to
    packed texts BY POSITION, so a renumbering that does not preserve the body's relative
    order makes every position mismatch — and the failure is silent, because a total mismatch
    just re-embeds the whole deck, which is exactly the old behaviour."""
    cid = _record(conn, sections=3)
    store.insert_chunks(conn, cid, ["A client remark."], kind="commentary",
                        sources=[{"kind": "comment", "author": "R. Vega"}])

    embeds = []
    real = embedding.embed
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (embeds.append(text),
                                                    real(text, timeout=timeout))[1])

    core.update_campaign(conn, campaign_id=cid, title="Bogota launch v2")

    assert len(embeds) <= 1, (
        f"a title edit re-embedded {len(embeds)} chunks — the position match has broken"
    )


# ---------------------------------------------------------------------------------------
# D109 — the paraphrase a lexical match cannot see


# ---------------------------------------------------------------------------------------
# D109 — attempted, measured, and REVERTED. It stays open.
#
# The row asks for a semantic fallback so a paraphrase sharing no vocabulary is noticed, and
# says it was blocked on D99 because `note` is a write path. D99 is fixed above, so the
# fallback was built — and then measured against the embedder this product actually ships:
#
#     one rule      0.651  0.771  0.872  0.880  0.896
#     different     0.378  0.391  0.418  0.446  0.495  0.695
#
# The populations overlap, so no threshold separates them. The row's own motivating pair
# scores 0.771 and a pair differing ONLY by negation scores 0.933 — a threshold catching the
# first fires on the second. What remains here is the LEXICAL path, unchanged in behaviour and
# now saying how it matched.


def test_a_resemblance_says_it_is_a_reading_and_how_it_was_reached(conn):
    """`basis: heuristic` because two rules being the same rule is a judgment somebody has to
    make, and `how` because "these share words" and "a model thinks these mean the same" are
    different claims — only one of them can be checked by looking. The second is not offered
    at all today, which is what D109 still records."""
    cid = core.ingest_campaign(conn, title="Bogota launch", market="Peru", status="concluded",
                               detail="A store opening.", confirm=True)["campaign_id"]
    corrections.note(conn, text="Creator captions name the product in the first line.",
                     campaign_id=cid, provenance="Client call, 3 March")

    second = corrections.note(conn,
                              text="Creator captions name the product on the first line.",
                              campaign_id=cid, provenance="Client call, 9 March")

    looks_like = second["new_correction"]["looks_like"]
    assert looks_like["basis"] == "heuristic"
    assert looks_like["how"] == "wording", (
        "the only matcher this product ships reads words, and must say so"
    )


def test_no_model_is_called_on_the_write_path(conn, monkeypatch):
    """The row's objection, which outlived its stated blocker. `note` is a WRITE — somebody
    is telling the library what a client said — and an attempt to add a semantic fallback here
    turned one write into one embed per correction on file, each with its own timeout: the
    "bounds when each call starts, not when the handler ends" pattern this very item fixes in
    `core`. Whatever closes D109, it is not that."""
    cid = core.ingest_campaign(conn, title="Bogota launch", market="Peru", status="concluded",
                               detail="A store opening.", confirm=True)["campaign_id"]
    for n in range(6):
        corrections.note(conn, text=f"A standing rule number {n} about seeding boxes.",
                         campaign_id=cid, provenance="Client call")

    calls = []
    real = embedding.embed
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (calls.append(text),
                                                    real(text, timeout=timeout))[1])

    corrections.note(conn, text="Something else entirely about packaging.",
                     campaign_id=cid, provenance="Client call")

    assert not calls, (
        f"recording one rule reached the embedder {len(calls)} times"
    )

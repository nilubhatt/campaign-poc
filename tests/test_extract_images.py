"""
Product feedback: image reuse detection can't require a separate manual upload per image -
"who'll use it then?" Decks are the primary path images enter the memory through, so
extraction has to pull embedded images out of the PPTX/PDF the user already uploads, not
require them uploaded again independently. This only applies when the actual file reaches
the server (asset_ref) - the LLM-first deck_text-only path never gives us bytes to extract
images from, an inherent limit, not a bug.
"""
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import config
import extract


def _make_image(path, seed=1, size=(64, 64)):
    rng = np.random.default_rng(seed)
    small = rng.integers(0, 256, size=(8, 8, 3), dtype="uint8")
    Image.fromarray(small, mode="RGB").resize(size, Image.BICUBIC).save(path)


def test_extract_images_from_pptx(tmp_path):
    img = tmp_path / "hero.png"
    _make_image(img)
    prs = Presentation()
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    found, warnings = extract.extract_images(deck)
    assert warnings == []
    assert len(found) == 1
    data, ext, slide_num = found[0]
    assert ext == ".png"
    assert len(data) > 0
    assert slide_num == 1


def test_extract_images_from_pptx_multiple_slides_distinct_images(tmp_path):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for seed in (1, 2, 3):
        img = tmp_path / f"img{seed}.png"
        _make_image(img, seed=seed)
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    found, _ = extract.extract_images(deck)
    assert len(found) == 3
    assert [slide_num for _, _, slide_num in found] == [1, 2, 3]


def test_extract_images_dedupes_the_same_image_repeated_across_slides(tmp_path):
    """A logo pasted on every slide shouldn't be stored/checked once per slide - and, more
    importantly, shouldn't consume the extraction cap and starve a later, distinct image
    (the actual demo risk: a hero image on a late slide never gets extracted)."""
    img = tmp_path / "logo.png"
    _make_image(img, seed=9)
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for _ in range(3):
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    found, _ = extract.extract_images(deck)
    assert len(found) == 1
    assert found[0][2] == 1  # kept the first occurrence


def test_extract_images_dedup_does_not_starve_a_later_distinct_image(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "MAX_EXTRACTED_IMAGES_PER_DECK", 2)
    logo = tmp_path / "logo.png"
    _make_image(logo, seed=9)
    hero = tmp_path / "hero.png"
    _make_image(hero, seed=42)

    prs = Presentation()
    layout = prs.slide_layouts[6]
    for _ in range(3):
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(str(logo), Inches(1), Inches(1), width=Inches(2))
    slide = prs.slides.add_slide(layout)
    slide.shapes.add_picture(str(hero), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    found, _ = extract.extract_images(deck)
    slide_nums = [slide_num for _, _, slide_num in found]
    assert 4 in slide_nums, "the distinct hero image on slide 4 must not be starved by the repeated logo"


def test_extract_images_from_pptx_picture_placeholder(tmp_path):
    """A picture inserted into a template's picture placeholder (common in corporate decks
    built from a layout) reports shape_type=PLACEHOLDER, not PICTURE - must still be found."""
    img = tmp_path / "hero.png"
    _make_image(img)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])  # "Picture with Caption"
    slide.placeholders[1].insert_picture(str(img))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    found, warnings = extract.extract_images(deck)
    assert len(found) == 1


def test_extract_images_from_pptx_grouped_picture(tmp_path):
    img = tmp_path / "hero.png"
    _make_image(img)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    p1 = slide.shapes.add_picture(str(img), Inches(0), Inches(0), width=Inches(1))
    p2_img = tmp_path / "hero2.png"
    _make_image(p2_img, seed=5)
    p2 = slide.shapes.add_picture(str(p2_img), Inches(2), Inches(0), width=Inches(1))
    slide.shapes.add_group_shape([p1, p2])
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    found, warnings = extract.extract_images(deck)
    assert len(found) == 2


def test_extract_images_from_pptx_one_bad_image_does_not_lose_the_rest(tmp_path, monkeypatch):
    """A single unsupported/corrupt embedded image (e.g. WebP - python-pptx's shape.image.ext
    raises for formats outside its known map) must be skipped, not abort extraction for the
    rest of the deck."""
    bad = tmp_path / "bad.png"
    _make_image(bad, seed=1)
    good = tmp_path / "hero.png"
    _make_image(good, seed=2)
    prs = Presentation()
    layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(layout)
    slide1.shapes.add_picture(str(bad), Inches(1), Inches(1), width=Inches(2))
    slide2 = prs.slides.add_slide(layout)
    slide2.shapes.add_picture(str(good), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    from pptx.shapes.picture import Picture
    real_ext = Picture.image.fget

    calls = {"n": 0}
    def flaky_image(self):
        calls["n"] += 1
        if calls["n"] == 1:
            raise ValueError("unrecognized image format")
        return real_ext(self)
    monkeypatch.setattr(Picture, "image", property(flaky_image))

    found, warnings = extract.extract_images(deck)
    assert warnings
    assert len(found) == 1
    assert found[0][2] == 2  # the good image, on slide 2, still came through


def test_extract_images_from_pptx_with_no_pictures(tmp_path):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    prs.slides.add_slide(layout)
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    found, warnings = extract.extract_images(deck)
    assert found == []
    assert warnings == []


def test_extract_images_from_pdf(tmp_path):
    img = tmp_path / "hero.png"
    _make_image(img)
    pdf_path = tmp_path / "deck.pdf"
    Image.open(img).convert("RGB").save(pdf_path, "PDF")

    found, warnings = extract.extract_images(pdf_path)
    assert warnings == []
    assert len(found) == 1
    data, ext, page_num = found[0]
    assert len(data) > 0
    assert page_num == 1


def test_extract_images_respects_max_cap(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "MAX_EXTRACTED_IMAGES_PER_DECK", 2)
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for seed in range(5):
        img = tmp_path / f"img{seed}.png"
        _make_image(img, seed=seed)
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    found, warnings = extract.extract_images(deck)
    assert len(found) == 2
    assert warnings  # truncation should be surfaced, not silent


def test_extract_images_unsupported_type_returns_empty_with_warning(tmp_path):
    found, warnings = extract.extract_images(Path("whatever.docx"), mime="application/msword")
    assert found == []
    assert warnings

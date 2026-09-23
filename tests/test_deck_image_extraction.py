"""
Product feedback: "Image reuse detection can't just be outside of the deck as an
independent functionality... you really want users to be doing this for every image? who'll
use it then?" upload_campaign now automatically extracts, fingerprints, embeds, and
reuse-checks images embedded in an uploaded deck (PPTX/PDF) when the actual file reaches the
server via asset_ref — no separate upload_image_asset call required. Only works via
asset_ref (a real file); the LLM-first deck_text-only path has no file to extract images
from, an inherent limit.
"""
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import core
import store


def _make_image(path, seed=1, size=(200, 200)):
    rng = np.random.default_rng(seed)
    small = rng.integers(0, 256, size=(8, 8, 3), dtype="uint8")
    Image.fromarray(small, mode="RGB").resize(size, Image.BICUBIC).save(path)


def _make_deck(path, img_path, slides=1):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for _ in range(slides):
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(str(img_path), Inches(1), Inches(1), width=Inches(2))
    prs.save(str(path))


def _make_multi_image_deck(path, img_paths):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for img_path in img_paths:
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(str(img_path), Inches(1), Inches(1), width=Inches(2))
    prs.save(str(path))


def test_upload_campaign_extracts_and_stores_deck_embedded_images(conn, tmp_path):
    img1 = tmp_path / "hero1.png"
    _make_image(img1, seed=1)
    img2 = tmp_path / "hero2.png"
    _make_image(img2, seed=2)
    deck = tmp_path / "deck.pptx"
    _make_multi_image_deck(deck, [img1, img2])

    result = core.ingest_campaign(conn, title="X", asset_ref={"path": str(deck)}, confirm=True)

    assert result["images_checked"] is True
    assert len(result["image_assets"]) == 2
    assert [e["location"] for e in result["image_assets"]] == [1, 2]
    for entry in result["image_assets"]:
        assert entry["fingerprinted"] is True
        assert entry["visually_embedded"] is True

    assets = store.get_assets_for_campaign(conn, result["campaign_id"])
    assert len(assets) == 2
    assert all(a["modality"] == "image" for a in assets)


def test_upload_campaign_dedupes_same_image_across_slides(conn, tmp_path):
    """The same image pasted on multiple slides (e.g. a logo) is stored/checked once, not
    once per slide - otherwise it fills the extraction cap and can starve a later, distinct
    image from ever being checked."""
    img = tmp_path / "logo.png"
    _make_image(img)
    deck = tmp_path / "deck.pptx"
    _make_deck(deck, img, slides=3)

    result = core.ingest_campaign(conn, title="X", asset_ref={"path": str(deck)}, confirm=True)

    assert len(result["image_assets"]) == 1
    assert result["image_assets"][0]["location"] == 1


def test_upload_campaign_with_no_embedded_images_has_empty_image_assets(conn, tmp_path):
    deck_no_images = tmp_path / "deck.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(deck_no_images))

    result = core.ingest_campaign(conn, title="X", asset_ref={"path": str(deck_no_images)}, confirm=True)
    assert result["image_assets"] == []


def test_upload_campaign_llm_first_path_has_no_image_assets(conn):
    """No file ever reaches the server on the deck_text-only path - nothing to extract from."""
    result = core.ingest_campaign(conn, title="X", deck_text="some text Claude already read",
                                  confirm=True)
    assert result["image_assets"] == []
    assert result["images_checked"] is False


def test_upload_campaign_extracts_images_when_asset_ref_given_alongside_deck_text(conn, tmp_path):
    """Docstring promises asset_ref works 'instead of/alongside deck_text'. This is also the
    realistic case: Claude reads the attached deck's text itself (deck_text) while the file
    also reaches the server via asset_ref (e.g. POST /upload) for storage - images must still
    be extracted, not silently skipped because deck_text happened to be present too."""
    img = tmp_path / "hero.png"
    _make_image(img)
    deck = tmp_path / "deck.pptx"
    _make_deck(deck, img, slides=1)

    result = core.ingest_campaign(conn, title="X", deck_text="text Claude already extracted",
                                  asset_ref={"path": str(deck)}, confirm=True)

    assert len(result["image_assets"]) == 1
    assert result["images_checked"] is True
    campaign = store.get_campaign(conn, result["campaign_id"])
    assert campaign["deck_text"] == "text Claude already extracted"  # caller's text wins, not overwritten


def test_upload_campaign_images_checked_true_even_with_no_images_found(conn, tmp_path):
    deck_no_images = tmp_path / "deck.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(deck_no_images))

    result = core.ingest_campaign(conn, title="X", asset_ref={"path": str(deck_no_images)}, confirm=True)
    assert result["image_assets"] == []
    assert result["images_checked"] is True


def test_upload_campaign_flags_reuse_from_a_deck_embedded_image(conn, tmp_path):
    """The actual scenario the feedback describes: the SAME hero image, reused in a deck
    uploaded for a different region, gets flagged automatically - no manual per-image step."""
    img = tmp_path / "hero.png"
    _make_image(img)

    mexico_deck = tmp_path / "mexico.pptx"
    _make_deck(mexico_deck, img)
    mexico = core.ingest_campaign(conn, title="Mexico Launch", region="LATAM",
                                  asset_ref={"path": str(mexico_deck)}, confirm=True)
    assert len(mexico["image_assets"]) == 1
    assert mexico["image_assets"][0]["reuse_flags"] == []  # nothing preceded it

    apac_deck = tmp_path / "apac.pptx"
    _make_deck(apac_deck, img)
    apac = core.ingest_campaign(conn, title="APAC Proposal", region="APAC",
                               asset_ref={"path": str(apac_deck)}, confirm=True)

    flags = apac["image_assets"][0]["reuse_flags"]
    assert len(flags) == 1
    assert flags[0]["campaign_id"] == mexico["campaign_id"]
    assert flags[0]["hamming_distance"] == 0
    assert "different region" in flags[0]["flag"]


def test_upload_campaign_reuse_flag_absent_when_same_region(conn, tmp_path):
    img = tmp_path / "hero.png"
    _make_image(img)

    deck1 = tmp_path / "a.pptx"
    _make_deck(deck1, img)
    core.ingest_campaign(conn, title="A", region="APAC", asset_ref={"path": str(deck1)}, confirm=True)

    deck2 = tmp_path / "b.pptx"
    _make_deck(deck2, img)
    result = core.ingest_campaign(conn, title="B", region="APAC", asset_ref={"path": str(deck2)}, confirm=True)

    assert result["image_assets"][0]["reuse_flags"][0]["flag"] is None


def test_upload_campaign_image_extraction_failure_does_not_block_the_campaign(conn, tmp_path, monkeypatch):
    """A corrupt/unextractable image shouldn't take down the whole upload - same
    per-step-failure philosophy as chunk embedding failures."""
    import extract
    img = tmp_path / "hero.png"
    _make_image(img)
    deck = tmp_path / "deck.pptx"
    _make_deck(deck, img)

    def boom(path, mime=None):
        raise RuntimeError("extraction blew up")
    monkeypatch.setattr(extract, "extract_images", boom)

    result = core.ingest_campaign(conn, title="X", asset_ref={"path": str(deck)}, confirm=True)
    assert result["campaign_id"]
    assert result["image_assets"] == []
    assert result["images_checked"] is False
    assert any("image" in w["detail"].lower() for w in result["warnings"])


def test_upload_campaign_legacy_ppt_reports_images_not_checked(conn, tmp_path):
    """Legacy .ppt is stored but has no image extractor - images_checked must say so
    explicitly rather than silently reporting an empty, clean-looking image_assets list."""
    fake_ppt = tmp_path / "old.ppt"
    fake_ppt.write_bytes(b"not a real ppt, just needs a resolvable path")

    result = core.ingest_campaign(conn, title="X", asset_ref={"path": str(fake_ppt)}, confirm=True)
    assert result["image_assets"] == []
    assert result["images_checked"] is False


def test_upload_campaign_fingerprint_persists_even_if_reuse_check_fails(conn, tmp_path, monkeypatch):
    """A failure computing reuse matches for THIS upload shouldn't discard the fingerprint
    that was already successfully stored - otherwise a later campaign's reuse check would
    never find this image, permanently, because of a transient failure now."""
    import core as core_module
    img = tmp_path / "hero.png"
    _make_image(img)
    deck = tmp_path / "deck.pptx"
    _make_deck(deck, img)

    def boom(*a, **k):
        raise RuntimeError("match query blew up")
    monkeypatch.setattr(core_module, "_phash_matches", boom)

    result = core.ingest_campaign(conn, title="X", asset_ref={"path": str(deck)}, confirm=True)
    entry = result["image_assets"][0]
    assert entry["fingerprinted"] is True
    assert entry["reuse_flags"] == []
    assert any(w["code"] == "reuse_check_failed" for w in result["warnings"])

    asset_id = entry["asset_id"]
    fp = store.get_all_fingerprints(conn)
    assert any(row["asset_id"] == asset_id for row in fp)


def test_upload_campaign_orphaned_image_file_cleaned_up_on_insert_failure(conn, tmp_path, monkeypatch):
    """If store.insert_asset fails after the image bytes were already written to disk, the
    orphaned file must be removed rather than leaked in ASSET_DIR forever."""
    import store as store_module
    img = tmp_path / "hero.png"
    _make_image(img)
    deck = tmp_path / "deck.pptx"
    _make_deck(deck, img)

    written_paths = []
    real_keep = core.config.ASSET_DIR

    def boom(conn, campaign_id, *, file_path, **kwargs):
        written_paths.append(real_keep / file_path)
        raise RuntimeError("db insert failed")
    monkeypatch.setattr(store_module, "insert_asset", boom)

    result = core.ingest_campaign(conn, title="X", asset_ref={"path": str(deck)}, confirm=True)
    assert result["image_assets"] == []
    assert written_paths and not written_paths[0].exists()

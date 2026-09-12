"""
Product review defect 05 (P1): "Failed embeddings leave unrepairable partial records."

The reviewer had two image assets stored and fingerprinted but carrying no visual embedding,
and wrote: "There is no tool that can backfill them — the only recovery is to upload the
images again." They also noted this was the second instance of the same class: text
embeddings that failed against Ollama saved campaigns with `embedded: false`, invisible to
every search while still listed. Chunking fixed that *trigger*; the unrepairable-state
*defect* was never addressed.

The acceptance line is "any partial record is recoverable without re-uploading source
material", and item 2.1 made that harder to dodge: a time budget now stops mid-deck on
purpose, so half-embedded records are a designed outcome rather than a rare accident. This
is the tool that finishes them.

Resumable by construction: reembed obeys the same budget as everything else, so on a large
backlog it does what it can, says what is left, and can simply be called again.
"""
import time

import pytest

import clip_embed
import config
import core
import embedding
import store


def _campaign_with_unembedded_chunks(conn, monkeypatch, count=6):
    """A campaign whose chunks exist but never got vectors - what a budget stop leaves."""
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.0)
    deck = "\n\n".join(f"section {i} " + "word " * 200 for i in range(count))
    created = core.ingest_campaign(conn, title="Interrupted", deck_text=deck, confirm=True)
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 45.0)
    return created


def test_finishes_a_campaign_the_budget_cut_short(conn, monkeypatch):
    created = _campaign_with_unembedded_chunks(conn, monkeypatch)
    assert created["chunks_embedded"] == 0, "precondition: nothing embedded yet"

    result = core.finish_indexing(conn)

    assert result["indexed"] > 0
    row = next(r for r in store.list_campaigns(conn) if r["id"] == created["campaign_id"])
    assert row["chunks_embedded"] == row["chunks_total"]
    assert row["embedded"] == 1, "a finished campaign reports itself searchable again"


def test_the_repaired_campaign_is_actually_findable(conn, monkeypatch):
    """The point is not the counter - it is that search stops missing the record."""
    created = _campaign_with_unembedded_chunks(conn, monkeypatch)
    assert core.find_similar(conn, text="section 3") == []

    core.finish_indexing(conn)

    hits = core.find_similar(conn, text="section 3")
    match = next(h for h in hits if h["campaign_id"] == created["campaign_id"])
    # Not just "a vector exists under this campaign" - the RIGHT chunk has to come back,
    # which a flag-only or constant-vector implementation would fail.
    assert "section 3" in match["matched_excerpt"]


def test_backfills_image_assets_that_never_got_a_visual_embedding(conn, tmp_path, monkeypatch):
    """The reviewer's literal situation: two assets stored and fingerprinted, no vector,
    and no way to fix them short of uploading the images again."""
    import numpy as np
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(3):
        img = tmp_path / f"i{i}.png"
        Image.fromarray(np.random.default_rng(i).integers(0, 256, (8, 8, 3), dtype="uint8"),
                        mode="RGB").resize((120, 120), Image.BICUBIC).save(img)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.0)
    created = core.ingest_campaign(conn, title="Images", asset_ref={"path": str(deck)},
                                   confirm=True)
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 45.0)
    assert created["images_embedded"] == 0

    result = core.finish_indexing(conn)

    assert result["images_indexed"] == 3
    assets = store.get_assets_for_campaign(conn, created["campaign_id"])
    assert all(a["embedded"] for a in assets)
    # The flags are not the point - visual search has to actually find them now, which an
    # implementation that skipped the "asset" vector space would fail.
    probe = config.ASSET_DIR / assets[0]["file_path"]
    similar = core.find_similar_images(conn, asset_ref={"path": str(probe)})
    assert similar["matches"], "the repaired assets must be visually searchable"


def test_a_nonexistent_campaign_id_is_an_error_not_a_quiet_success(conn):
    """Reporting "complete, nothing to do" for a typo'd id tells the caller the record is
    fine when it may not even exist. Every sibling tool says "not found"."""
    with pytest.raises(ValueError, match="not found"):
        core.finish_indexing(conn, campaign_id="camp_does_not_exist")

    with pytest.raises(ValueError):
        core.finish_indexing(conn, campaign_id="")


def test_an_untouched_campaign_is_not_marked_as_modified(conn, monkeypatch):
    """A global run refreshed every campaign's rollup, rewriting updated_at on records it
    never touched - a modification timestamp for an edit that never happened."""
    clean = core.ingest_campaign(conn, title="Already fine", deck_text="all good",
                                 confirm=True)
    before = store.get_campaign(conn, clean["campaign_id"])["updated_at"]

    _campaign_with_unembedded_chunks(conn, monkeypatch, count=3)
    core.finish_indexing(conn)

    assert store.get_campaign(conn, clean["campaign_id"])["updated_at"] == before


def test_can_be_scoped_to_one_campaign(conn, monkeypatch):
    first = _campaign_with_unembedded_chunks(conn, monkeypatch, count=3)
    second = _campaign_with_unembedded_chunks(conn, monkeypatch, count=3)

    core.finish_indexing(conn, campaign_id=first["campaign_id"])

    rows = {r["id"]: r for r in store.list_campaigns(conn)}
    assert rows[first["campaign_id"]]["chunks_embedded"] > 0
    assert rows[second["campaign_id"]]["chunks_embedded"] == 0, "other campaigns untouched"


def test_reports_nothing_to_do_when_everything_is_embedded(conn):
    core.ingest_campaign(conn, title="Fine", deck_text="a brief", confirm=True)

    result = core.finish_indexing(conn)

    assert result["indexed"] == 0
    assert result["remaining"] == 0
    assert result["complete"] is True


def test_obeys_the_same_time_budget_and_says_what_is_left(conn, monkeypatch):
    """A backlog can be larger than any one call. It must stop cleanly and be callable
    again rather than outliving the transport - the failure this whole phase is about."""
    created = _campaign_with_unembedded_chunks(conn, monkeypatch, count=12)

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.25)
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (time.sleep(0.1),
                                                    [0.0] * config.EMBED_DIM)[1])

    started = time.monotonic()
    result = core.finish_indexing(conn)
    elapsed = time.monotonic() - started

    assert elapsed < 1.0
    assert result["complete"] is False
    assert result["remaining"] > 0
    assert "again" in result["note"].lower(), "must tell the caller it can resume"
    assert str(result["remaining"]) in result["note"]


def test_calling_it_repeatedly_finishes_the_backlog(conn, monkeypatch):
    """Resumability is the whole design: two modest calls must achieve what one big one
    could not."""
    created = _campaign_with_unembedded_chunks(conn, monkeypatch, count=8)

    for _ in range(10):
        if core.finish_indexing(conn)["complete"]:
            break

    row = next(r for r in store.list_campaigns(conn) if r["id"] == created["campaign_id"])
    assert row["chunks_embedded"] == row["chunks_total"] > 0


def test_a_failing_embedder_does_not_lose_the_rest_of_the_backlog(conn, monkeypatch):
    created = _campaign_with_unembedded_chunks(conn, monkeypatch, count=4)

    calls = {"n": 0}
    real = embedding.embed

    def flaky(text, timeout=None):
        calls["n"] += 1
        if calls["n"] == 1:
            raise ValueError("embedder hiccup")
        return real(text, timeout)

    monkeypatch.setattr(embedding, "embed", flaky)

    result = core.finish_indexing(conn)

    assert result["indexed"] > 0, "one bad chunk must not abort the run"
    assert result["errors"], "and the failure is reported, not swallowed"


# ── the failure modes the happy path hides (design review) ──────────────────

def test_a_dead_embedder_does_not_tell_the_caller_to_keep_retrying(conn, monkeypatch):
    """With Ollama down every item fails in milliseconds, so the run "finishes" instantly
    with nothing done and a full backlog. Saying "call again to continue" there is an
    invitation to loop forever against a component that is not coming back on its own."""
    _campaign_with_unembedded_chunks(conn, monkeypatch, count=6)

    import httpx
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (_ for _ in ()).throw(
                            ValueError("could not reach the ollama embedder")))

    result = core.finish_indexing(conn)

    assert result["indexed"] == 0
    note = (result.get("note") or "").lower()
    assert "will not help" in note, (
        "no progress means retrying changes nothing - say what is actually wrong"
    )
    assert "resumes where it stopped" not in note, "must not invite a retry loop"
    assert result["complete"] is False


def test_repeated_identical_failures_are_collapsed_not_repeated(conn, monkeypatch):
    """200 copies of the same Ollama message is not a report."""
    _campaign_with_unembedded_chunks(conn, monkeypatch, count=8)
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (_ for _ in ()).throw(
                            ValueError("could not reach the ollama embedder at localhost")))

    result = core.finish_indexing(conn)

    assert len(result["errors"]) == 1, f"one line per distinct cause: {result['errors']}"
    assert result["errors"][0]["count"] >= 6
    assert "ollama" in result["errors"][0]["reason"].lower()


def test_items_that_can_never_succeed_do_not_block_completion(conn, tmp_path, monkeypatch):
    """An asset whose file was deleted fails identically on every run. Counting it as
    "remaining" means `complete` never becomes true and the caller loops forever."""
    import numpy as np
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    img = tmp_path / "i.png"
    Image.fromarray(np.random.default_rng(1).integers(0, 256, (8, 8, 3), dtype="uint8"),
                    mode="RGB").resize((120, 120), Image.BICUBIC).save(img)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.0)
    created = core.ingest_campaign(conn, title="Gone", asset_ref={"path": str(deck)},
                                   confirm=True)
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 45.0)

    for asset in store.get_assets_for_campaign(conn, created["campaign_id"]):
        (config.ASSET_DIR / asset["file_path"]).unlink()

    result = core.finish_indexing(conn)

    assert result["failed"] >= 1, "unfixable items are reported as failed, not pending"
    assert result["complete"] is True, "nothing is left that another call could fix"


def test_the_newest_backlog_is_finished_first(conn, monkeypatch):
    """The deck the marketer just uploaded is the one they are waiting on; finishing an
    eight-month-old backlog first is the wrong order."""
    old = _campaign_with_unembedded_chunks(conn, monkeypatch, count=4)
    time.sleep(0.01)
    new = _campaign_with_unembedded_chunks(conn, monkeypatch, count=4)

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.25)
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (time.sleep(0.1),
                                                    [0.0] * config.EMBED_DIM)[1])

    core.finish_indexing(conn)

    rows = {r["id"]: r for r in store.list_campaigns(conn)}
    assert rows[new["campaign_id"]]["chunks_embedded"] > 0, "newest first"
    assert rows[old["campaign_id"]]["chunks_embedded"] == 0


def test_search_says_when_results_may_be_incomplete(conn, monkeypatch):
    """The review's second complaint was that a half-indexed record is invisible to search
    while looking fine elsewhere. Listing it honestly is half the fix; the moment it
    actually misleads someone is when they ask a question and get a confident answer."""
    core.ingest_campaign(conn, title="Indexed", deck_text="mexico launch activity",
                         confirm=True)
    _campaign_with_unembedded_chunks(conn, monkeypatch, count=4)

    result = core.find_similar_with_context(conn, text="mexico launch")

    assert result["matches"]
    assert any("partly searchable" in w.lower() or "incomplete" in w.lower()
               for w in result["warnings"]), result["warnings"]

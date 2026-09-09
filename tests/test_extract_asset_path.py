"""Exercises the asset_ref (file) ingestion path with a real .pptx, not just LLM-first
flat deck_text — makes sure extract.extract_units() feeds natural per-slide chunks."""
from pptx import Presentation

import core
import extract


def _make_pptx(path, slide_texts):
    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content
    for text in slide_texts:
        slide = prs.slides.add_slide(layout)
        slide.shapes.placeholders[1].text_frame.text = text
    prs.save(str(path))


def test_extract_units_returns_one_unit_per_slide(tmp_path):
    pptx_path = tmp_path / "deck.pptx"
    _make_pptx(pptx_path, ["First slide content", "Second slide content", "Third slide content"])

    units, warnings = extract.extract_units(pptx_path)

    assert warnings == []
    assert len(units) == 3
    assert units[0] == "First slide content"
    assert units[2] == "Third slide content"


def test_ingest_via_asset_ref_uses_natural_slide_units(conn, tmp_path):
    pptx_path = tmp_path / "deck.pptx"
    _make_pptx(pptx_path, ["Audience: young adults", "Channel: social + OOH", "Budget: 500k"])

    r = core.ingest_campaign(conn, title="From Deck", status="concluded",
                             asset_ref={"path": str(pptx_path)})

    assert r["chunks_total"] >= 1  # small slides may merge into one chunk via pack()
    assert r["chunks_embedded"] == r["chunks_total"]
    assert r["warnings"] == []

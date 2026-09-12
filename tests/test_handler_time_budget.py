"""
Product review defect 04 (P1): "Model init inside the request handler exceeds the MCP
timeout" — and the instruction attached to it: "Audit every handler for lazy initialisation
— this is a pattern, not a single site."

Item 1.1 fixed the site the reviewer hit (the installed binary never called warm_up, so the
vision model loaded inside the first image call). This item is the sweep, and it found a
second site that is likelier to bite in practice:

`embedding.embed()` had `timeout=60` — the SAME as the transport ceiling — and
`ingest_campaign` calls it once per chunk. A twelve-chunk deck against a stalled Ollama
could block for twelve minutes inside one tool call while the client gave up at sixty
seconds. The user sees "did not respond within 60s", the work is half-done, and nothing
says which half.

So: a per-call timeout well under the ceiling, a wall-clock budget across the whole loop,
and a partial result that says what it managed and how to finish it.
"""
import time

import pytest

import config
import core
import embedding
import store


def test_per_call_embed_timeout_is_well_under_the_transport_ceiling():
    """Equal to the ceiling is useless: the handler has other work to do, and usually makes
    more than one call."""
    assert config.EMBED_TIMEOUT_SECONDS < config.TOOL_TIME_BUDGET_SECONDS / 2, (
        "a single embed call must not be able to consume the whole tool budget"
    )


def test_a_slow_embedder_does_not_run_past_the_budget(conn, monkeypatch):
    """The behaviour that matters: stop and report, rather than block until the client
    disconnects and leaves the user with a timeout and no explanation."""
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.3)

    calls = {"n": 0}

    def slow_embed(text, timeout=None):
        calls["n"] += 1
        time.sleep(0.15)
        return [0.0] * config.EMBED_DIM

    monkeypatch.setattr(embedding, "embed", slow_embed)

    long_deck = "\n\n".join(f"section {i} " + "word " * 200 for i in range(12))
    started = time.monotonic()
    result = core.ingest_campaign(conn, title="Slow", deck_text=long_deck, confirm=True)
    elapsed = time.monotonic() - started

    # 12 chunks x 0.15s = 1.8s if it ground through them all, so a 2.0s bound could not
    # fail. Bound it just above the budget instead.
    assert elapsed < 0.9, "must abandon the loop, not grind through every chunk"
    assert calls["n"] <= 3, "should have stopped within a call or two of the budget"
    assert result["chunks_embedded"] < result["chunks_total"]


def test_running_out_of_budget_says_what_happened_and_how_to_finish(conn, monkeypatch):
    """Half-embedded records were the review's other complaint (defect 05): silently
    invisible to search while still appearing in list_campaigns. If we stop early we have to
    say so, in words an operator can act on."""
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.2)
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (time.sleep(0.12), [0.0] * config.EMBED_DIM)[1])

    long_deck = "\n\n".join(f"section {i} " + "word " * 200 for i in range(10))
    result = core.ingest_campaign(conn, title="Slow", deck_text=long_deck, confirm=True)

    budget_warnings = [w for w in result["warnings"] if "budget" in w.lower()]
    assert budget_warnings, f"no warning explained the early stop: {result['warnings']}"
    warning = budget_warnings[0]
    assert str(result["chunks_embedded"]) in warning and str(result["chunks_total"]) in warning
    assert "saved" in warning.lower(), "must say the upload itself survived"
    # 2.1 deliberately left this open: naming a recovery tool that did not exist yet would
    # have sent a marketer after something Claude could not find. 2.2 built it, so the
    # sentence can now be closed - this is that promise being kept.
    assert "finish_indexing" in warning, "must name the tool that finishes the job"
    assert "no re-upload" in warning.lower()


def test_a_fast_embedder_still_completes_everything(conn):
    """The budget must not truncate normal work - only rescue a pathological one."""
    deck = "\n\n".join(f"section {i} " + "word " * 50 for i in range(8))
    result = core.ingest_campaign(conn, title="Fast", deck_text=deck, confirm=True)

    assert result["chunks_embedded"] == result["chunks_total"] > 0
    assert not [w for w in result["warnings"] if "budget" in w.lower()]


def test_the_campaign_survives_even_if_nothing_embeds(conn, monkeypatch):
    """Storage and search are separate concerns; losing the embedder must not lose the
    upload."""
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.0)

    result = core.ingest_campaign(conn, title="Budget gone", deck_text="a b c", confirm=True)

    assert result["campaign_id"]
    assert store.get_campaign(conn, result["campaign_id"]) is not None
    assert result["embedded"] is False


def test_the_image_loop_shares_one_budget_with_the_text_loop(conn, tmp_path, monkeypatch):
    """Both loops live in the same tool call, so they must budget against the same clock.
    Two separate budgets would let a handler take twice the ceiling and still believe it
    was within it - and the image loop runs first, so it silently spends what the text loop
    then measures against."""
    import numpy as np
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    import clip_embed

    deck = tmp_path / "deck.pptx"
    prs = Presentation()
    for i in range(6):
        img = tmp_path / f"i{i}.png"
        rng = np.random.default_rng(i)
        Image.fromarray(rng.integers(0, 256, (8, 8, 3), dtype="uint8"), mode="RGB").resize(
            (120, 120), Image.BICUBIC).save(img)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    prs.save(str(deck))

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.25)
    monkeypatch.setattr(clip_embed, "embed_image",
                        lambda p: (time.sleep(0.1), [0.0] * config.CLIP_EMBED_DIM)[1])

    started = time.monotonic()
    result = core.ingest_campaign(conn, title="Heavy deck",
                                  asset_ref={"path": str(deck)}, confirm=True)
    elapsed = time.monotonic() - started

    assert elapsed < 1.2, "the image loop must respect the budget too"
    assert any("budget" in w.lower() for w in result["warnings"])
    assert result["campaign_id"], "the campaign is still saved"


def test_no_single_call_may_run_past_the_deadline(conn, monkeypatch):
    """Checking the clock at the top of an iteration bounds when work STARTS, not when it
    ends - so the last call could begin a millisecond inside the budget and then run for a
    full per-call timeout on top. With the shipped defaults that arithmetic came to exactly
    the transport ceiling (45 + 15 = 60), i.e. a budget that can still be cut off, which is
    no budget at all. A call must be given only the time that is actually left."""
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.4)
    monkeypatch.setattr(config, "EMBED_TIMEOUT_SECONDS", 5.0)

    granted = []

    def record_timeout(text, timeout=None):
        granted.append(timeout)
        return [0.0] * config.EMBED_DIM

    monkeypatch.setattr(embedding, "embed", record_timeout)

    deck = "\n\n".join(f"section {i} " + "word " * 200 for i in range(6))
    core.ingest_campaign(conn, title="Bounded", deck_text=deck, confirm=True)

    assert granted, "the embedder should have been called"
    assert all(t is not None for t in granted), "every call must be given a deadline"
    assert all(t <= config.EMBED_TIMEOUT_SECONDS for t in granted)
    assert all(t <= 0.4 for t in granted), (
        "a call may never be granted more time than the budget has left - otherwise the "
        "handler's worst case is budget + timeout, which is what blew the ceiling"
    )


def _deck_with_images(tmp_path, count):
    import numpy as np
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(count):
        img = tmp_path / f"img{i}.png"
        rng = np.random.default_rng(i)
        Image.fromarray(rng.integers(0, 256, (8, 8, 3), dtype="uint8"), mode="RGB").resize(
            (120, 120), Image.BICUBIC).save(img)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))
    return deck


def test_reuse_checking_is_never_cut_by_the_budget(conn, tmp_path, monkeypatch):
    """Reuse detection is the question this product exists to answer, and it is cheap -
    storing a file and hashing it is milliseconds. Only the expensive visual embedding is
    worth abandoning. So every image gets stored and reuse-checked even when the budget is
    already blown; the CLIP pass is what yields."""
    import clip_embed

    deck = _deck_with_images(tmp_path, 5)
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.0)   # no budget at all
    monkeypatch.setattr(clip_embed, "embed_image",
                        lambda p: pytest.fail("must not embed with no budget left"))

    result = core.ingest_campaign(conn, title="Deck", asset_ref={"path": str(deck)},
                                  confirm=True)

    assert len(result["image_assets"]) == 5, "every image must still be stored"
    assert all(e["fingerprinted"] for e in result["image_assets"])
    assert not any(e["visually_embedded"] for e in result["image_assets"])
    assert result["images_checked"] is True, "reuse WAS checked for all of them"


def test_a_budget_stop_never_claims_images_were_checked_when_they_were_not(conn, tmp_path, monkeypatch):
    """The docstring promises Claude that images_checked=True means reuse was actually
    checked, and that it must not say "no reuse found" when False. If the loop stops early
    and still reports True, Claude tells a marketer there is no reuse in slides nobody
    looked at - which is the stored-versus-checked confusion this whole review opened with."""
    import extract

    deck = _deck_with_images(tmp_path, 4)
    real_extract = extract.extract_images

    def half_extract(path, mime=None):
        found, warns = real_extract(path, mime)
        raise RuntimeError("extraction died halfway")

    monkeypatch.setattr(extract, "extract_images", half_extract)

    result = core.ingest_campaign(conn, title="Deck", asset_ref={"path": str(deck)},
                                  confirm=True)

    assert result["images_checked"] is False
    assert result["image_assets"] == []


def test_the_response_says_how_many_images_were_embedded_of_how_many(conn, tmp_path, monkeypatch):
    """Counts, not a boolean: "stored" and "searchable" must never be conflated (defect 05)."""
    import clip_embed

    deck = _deck_with_images(tmp_path, 4)
    calls = {"n": 0}

    def one_then_stop(path):
        calls["n"] += 1
        if calls["n"] > 1:
            time.sleep(0.3)
        return [0.0] * config.CLIP_EMBED_DIM

    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.2)
    monkeypatch.setattr(clip_embed, "embed_image", one_then_stop)

    result = core.ingest_campaign(conn, title="Deck", asset_ref={"path": str(deck)},
                                  confirm=True)

    assert result["images_total"] == 4
    assert 0 < result["images_embedded"] < 4
    assert result["images_checked"] is True, "all four were still reuse-checked"


def test_every_unembedded_image_left_a_row_that_can_be_finished_later(conn, tmp_path, monkeypatch):
    """The asymmetry review flagged: a text chunk that misses the budget leaves a row a
    later pass can find, but an image that was never extracted leaves nothing at all - so
    the promised remedy could never work on it. Both modalities must leave a row."""
    import clip_embed

    deck = _deck_with_images(tmp_path, 4)
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.0)
    monkeypatch.setattr(clip_embed, "embed_image",
                        lambda p: pytest.fail("no budget"))

    result = core.ingest_campaign(conn, title="Deck", asset_ref={"path": str(deck)},
                                  confirm=True)

    assets = store.get_assets_for_campaign(conn, result["campaign_id"])
    assert len(assets) == 4, "all four rows exist and can be finished later"
    assert all(a["embedded"] == 0 for a in assets)


# ── the sweep's own blind spot: the text embedder was never warmed ───────────

def test_the_text_embedder_is_warmed_at_startup(monkeypatch):
    """This item exists because first-use initialisation inside a handler blows the
    transport ceiling - and the sweep nearly shipped with the text embedder doing exactly
    that. CLIP was warmed at startup; Ollama was not, so it unloads the model after its idle
    period and the first chunk of every upload pays the reload inside the handler."""
    import embedding

    called = {"n": 0}
    monkeypatch.setattr(config, "EMBED_PROVIDER", "ollama")
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: called.__setitem__("n", called["n"] + 1))

    embedding.warm_up()

    assert called["n"] == 1, "startup must make one small embed call to load the model"


def test_warming_the_embedder_never_takes_the_server_down(monkeypatch):
    """Same rule as the vision model: text search failing to warm must not stop a server
    whose other tools work fine."""
    import embedding

    monkeypatch.setattr(config, "EMBED_PROVIDER", "ollama")
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: (_ for _ in ()).throw(
                            RuntimeError("ollama is not running")))

    embedding.warm_up()   # must not raise


def test_the_offline_hash_provider_needs_no_warming(monkeypatch):
    import embedding

    called = {"n": 0}
    monkeypatch.setattr(config, "EMBED_PROVIDER", "hash")
    monkeypatch.setattr(embedding, "embed",
                        lambda text, timeout=None: called.__setitem__("n", called["n"] + 1))

    embedding.warm_up()

    assert called["n"] == 0


def test_ollama_is_asked_to_keep_the_model_resident(monkeypatch):
    """Warming once is pointless if Ollama unloads the model minutes later and the next
    upload pays the load again."""
    import embedding

    sent = {}

    class _Resp:
        def raise_for_status(self):
            pass

        def json(self):
            return {"embedding": [0.0] * config.EMBED_DIM}

    def fake_post(url, json=None, timeout=None):
        sent.update(json or {})
        return _Resp()

    import httpx
    monkeypatch.setattr(httpx, "post", fake_post)
    embedding._embed_ollama("hello", 5.0)

    assert "keep_alive" in sent, "ask Ollama to keep the embedding model loaded"

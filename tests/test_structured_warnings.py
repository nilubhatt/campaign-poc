"""
Product review defect 09 (P2): "Warnings are engineer-readable, not operator-actionable."

The reviewer was generous about the engineering and precise about the gap:

    "The download-failure warning was precise enough to diagnose from — genuinely good
    engineering output. But the intended user is a marketer, for whom 'Failed to download
    weights for tag openai' carries no action at all."

    Fix: "Give warnings a structured shape — a stable `code`, a one-line human `remedy`, and
    the raw `detail` — so the surface can say 'Visual search is offline — ask IT to run
    setup' while the detail stays available for support."

Three audiences, one warning: the marketer needs the remedy, support needs the detail, and
anything programmatic needs the code — which is the part a reworded message must not break.
"""
import pytest
from pptx import Presentation

import clip_embed
import config
import core
import extract
import notices
import store


def test_a_warning_carries_a_code_a_remedy_and_the_raw_detail(conn):
    result = core.ingest_campaign(conn, title="   ", confirm=True)

    warning = result["warnings"][0]
    assert warning["code"], "stable identifier — the part a reworded message must not break"
    assert warning["remedy"], "one line the marketer can act on"
    assert warning["detail"], "what support needs, kept rather than replaced"


def test_the_remedy_is_written_for_a_marketer_not_an_engineer(monkeypatch, tmp_path, conn):
    """The review's own example. "Failed to download weights for tag 'openai'" is a fine
    diagnostic and a useless instruction."""
    image = tmp_path / "hero.png"
    _write_png(image)

    def unavailable(*a, **k):
        raise RuntimeError("Failed to download weights for tag 'openai': connection refused")

    monkeypatch.setattr(clip_embed, "embed_image", unavailable)
    # The machine the reviewer ran on: the model is not there at all, which is the
    # operator's problem to fix once — as opposed to one bad PNG on a working model.
    monkeypatch.setattr(clip_embed, "weights_status", lambda: clip_embed.WeightsResolution(
        ok=False, source="missing", reason="no checkpoint found"))

    result = core.ingest_image_asset(conn, campaign_id=_campaign(conn),
                                    asset_ref={"path": str(image)})

    warning = next(w for w in result["warnings"] if w["code"] == "visual_search_offline")
    assert "weights" not in warning["remedy"].lower(), "a marketer does not know what a weight is"
    assert "openai" not in warning["remedy"].lower()
    assert "Failed to download weights" in warning["detail"], "support still gets the real text"


def test_every_code_in_use_has_a_remedy_written_for_it(conn):
    """A code with no remedy is a warning that reverted to being engineer-only, which is the
    defect. The registry is what stops one being added quietly."""
    for code in notices.CODES:
        assert notices.remedy_for(code), code


def test_an_unregistered_code_is_a_programming_error_not_a_blank_remedy(conn):
    with pytest.raises(KeyError):
        notices.notice("no_such_code", detail="whatever")


def test_the_code_survives_a_reworded_message(conn):
    """The point of a stable code: support tooling and the installer's own checks key off it,
    so improving the wording must not be a breaking change."""
    first = notices.notice("nothing_to_embed", detail="no title/detail/deck_text")

    assert first["code"] == "nothing_to_embed"
    assert first["detail"] == "no title/detail/deck_text"


def test_warnings_of_the_same_kind_are_not_repeated_once_per_item(tmp_path, conn, monkeypatch):
    """A deck with twenty unreadable images produced twenty near-identical lines, which is
    how a genuinely important warning gets scrolled past. One entry, with the count."""
    deck = _deck_with_images(tmp_path, 6)
    monkeypatch.setattr(core.images, "phash",
                        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("bad image")))

    result = core.ingest_campaign(conn, title="Broken images",
                                  asset_ref={"path": str(deck)}, confirm=True)

    repeated = [w for w in result["warnings"] if w["code"] == "image_not_fingerprinted"]
    assert len(repeated) == 1, f"{len(repeated)} near-identical lines"
    assert repeated[0]["count"] == 6
    assert "bad image" in repeated[0]["detail"], "the first real reason is still there"


def test_the_budget_warning_keeps_its_instruction(tmp_path, conn, monkeypatch):
    """This one was already operator-actionable and says exactly what to do next; the
    restructuring must not lose it."""
    monkeypatch.setattr(config, "TOOL_TIME_BUDGET_SECONDS", 0.0)
    deck = "\n\n".join(f"section {i} " + "word " * 200 for i in range(6))

    result = core.ingest_campaign(conn, title="Cut short", deck_text=deck, confirm=True)

    warning = next(w for w in result["warnings"] if w["code"] == "indexing_incomplete")
    assert "finish_indexing" in warning["remedy"]
    assert result["campaign_id"] in warning["remedy"], "the call has to be usable as written"


def test_an_extraction_warning_is_structured_too(tmp_path):
    """extract.py raises warnings that reach the same response, so a caller cannot be left
    handling two shapes depending on which layer produced the line."""
    import zipfile

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "broken.pptx"
    prs.save(str(path))
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("ppt/modernComments/modernComment_1.xml", "<not-xml")

    _, warnings = extract.extract_commentary(path)

    assert warnings[0]["code"] == "commentary_unreadable"
    assert warnings[0]["remedy"]


def test_a_severity_says_whether_the_user_has_to_do_anything(conn):
    """"Visual search is offline" and "one image out of forty was skipped" are not the same
    news, and a flat list made a surface treat them identically."""
    result = core.ingest_campaign(conn, title="   ", confirm=True)

    assert result["warnings"][0]["severity"] in notices.SEVERITIES


def test_the_response_says_what_the_user_should_be_told(tmp_path, conn, monkeypatch):
    """With several warnings, a surface needs to know which one leads — otherwise it reads
    them all out, which is what a marketer ignores."""
    image = tmp_path / "hero.png"
    _write_png(image)
    monkeypatch.setattr(clip_embed, "embed_image",
                        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no weights")))
    monkeypatch.setattr(clip_embed, "weights_status", lambda: clip_embed.WeightsResolution(
        ok=False, source="missing", reason="no checkpoint found"))

    result = core.ingest_image_asset(conn, campaign_id=_campaign(conn),
                                    asset_ref={"path": str(image)})

    assert result["warnings"], "precondition"
    leading = notices.leading(result["warnings"])
    assert leading["code"] == "visual_search_offline"


# ── helpers ─────────────────────────────────────────────────────────────────

def _campaign(conn):
    return core.ingest_campaign(conn, title="Host", detail="a campaign", confirm=True)[
        "campaign_id"]


def _write_png(path, seed=0):
    import numpy as np
    from PIL import Image

    data = np.random.default_rng(seed).integers(0, 256, (8, 8, 3), dtype="uint8")
    Image.fromarray(data, mode="RGB").resize((120, 120), Image.BICUBIC).save(path)
    return path


def _deck_with_images(tmp_path, n):
    from pptx.util import Inches

    prs = Presentation()
    for i in range(n):
        img = _write_png(tmp_path / f"i{i}.png", seed=i)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(img), Inches(1), Inches(1), width=Inches(2))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))
    return deck
